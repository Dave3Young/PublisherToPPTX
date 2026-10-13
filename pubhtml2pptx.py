#!/usr/bin/env python3
"""
pubhtml2pptx.py
===============

Converts the filtered-HTML + PNG output of PublisherPubToHtmlForPptx.ps1
into a .pptx whose slides reproduce the original Publisher page layout, so the
result can be imported into Canva (or PowerPoint/Google Slides) with the text
boxes and pictures arriving as SEPARATE EDITABLE ELEMENTS rather than one
flattened page image.

Expected input (exactly what the PowerShell export script produces):

    <source folder>\\
        Example.pub
        Example\\                                  <-- the export subfolder
            Example.htm                            <-- filtered HTML
            Example_files\\                        <-- Publisher's web assets
            Example_page_001_image_001.png         <-- 300dpi COM exports
            Example_masterpage_001_image_004.png
            Example_scratch_000_image_007.png
            Example_scratch_000_text_008.png
            Example_scratch_text.txt

Usage
-----
    # one export folder (or point straight at the .htm)
    python pubhtml2pptx.py "C:\\Pubs\\Example"

    # every export folder beneath a tree
    python pubhtml2pptx.py "C:\\Pubs" --recurse

    # use the 300dpi PNGs in place of Publisher's web-quality JPEGs
    python pubhtml2pptx.py "C:\\Pubs" --recurse --hires

Options
-------
    -o, --output PATH   output .pptx (single input) or output folder (--recurse)
    --recurse           find every *_export folder beneath the input folder
    --dpi N             CSS pixels per inch, default 96
    --hires             swap web images for the matching 300dpi COM PNG
    --no-extras         do not append slides for scratch-area PNGs/text
    --reflow            let text wrap freely instead of breaking lines where
                        Publisher does
    --page-size WxH     fallback page size in inches when stage 1 recorded none,
                        default 8.5x11
    --report            write <basename>_conversion_report.txt beside the pptx
    -v, --verbose       per-shape logging

Disclaimer: provided as is, without warranty. Test on copies.
Written for the Publisher retirement toolkit at www.david-e-young.com.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import struct
import sys
import glob as globmod
import dataclasses
from dataclasses import dataclass, field
from typing import Optional
from urllib.parse import unquote

try:
    from bs4 import BeautifulSoup, Comment, NavigableString, Tag
    # Base of Comment, Doctype, CData etc.: strings that are not page text.
    # Publisher 2010 puts its VML shape markup in <!--[if gte vml 1]> comments.
    from bs4.element import PreformattedString
except ImportError:
    sys.exit("Missing dependency. Run:  pip install beautifulsoup4 lxml python-pptx Pillow")

try:
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.oxml.ns import qn
except ImportError:
    sys.exit("Missing dependency. Run:  pip install python-pptx")

try:
    from PIL import Image
    HAVE_PIL = True
except ImportError:
    HAVE_PIL = False

EMU_PER_INCH = 914400
PPTX_MAX_IN = 56.0

# Outline shapes drawn behind rebuilt text boxes
SHAPE_GEOMS = {
    "rect": MSO_SHAPE.RECTANGLE,
    "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "wedgeRoundRectCallout": MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT,
    "wedgeRectCallout": MSO_SHAPE.RECTANGULAR_CALLOUT,
    "wedgeEllipseCallout": MSO_SHAPE.OVAL_CALLOUT,
}
# VML dashstyle -> PowerPoint dash. python-pptx writes ROUND_DOT as sysDot (1:1 dots)
# and SQUARE_DOT as sysDash (3:1 short dashes), so dots map to ROUND_DOT.
LINE_DASHES = {
    "dot": MSO_LINE_DASH_STYLE.ROUND_DOT, "1 1": MSO_LINE_DASH_STYLE.ROUND_DOT,
    "shortdot": MSO_LINE_DASH_STYLE.ROUND_DOT, "dash": MSO_LINE_DASH_STYLE.DASH,
    "shortdash": MSO_LINE_DASH_STYLE.SQUARE_DOT, "dashdot": MSO_LINE_DASH_STYLE.DASH_DOT,
    "shortdashdot": MSO_LINE_DASH_STYLE.DASH_DOT, "longdash": MSO_LINE_DASH_STYLE.LONG_DASH,
    "longdashdot": MSO_LINE_DASH_STYLE.LONG_DASH_DOT,
    "longdashdotdot": MSO_LINE_DASH_STYLE.DASH_DOT_DOT,
    "shortdashdotdot": MSO_LINE_DASH_STYLE.DASH_DOT_DOT,
}

# ----------------------------------------------------------------------------
# Small CSS engine
# ----------------------------------------------------------------------------

INHERITED_PROPS = {
    "font-family", "font-size", "font-weight", "font-style", "font-variant",
    "color", "text-align", "line-height", "letter-spacing", "text-indent",
    "text-transform", "direction", "white-space",
    "-x-script",                 # ours: "super" or "sub", from <sup>/<sub>
}

NAMED_COLORS = {
    "black": "000000", "silver": "C0C0C0", "gray": "808080", "grey": "808080",
    "white": "FFFFFF", "maroon": "800000", "red": "FF0000", "purple": "800080",
    "fuchsia": "FF00FF", "magenta": "FF00FF", "green": "008000", "lime": "00FF00",
    "olive": "808000", "yellow": "FFFF00", "navy": "000080", "blue": "0000FF",
    "teal": "008080", "aqua": "00FFFF", "cyan": "00FFFF", "orange": "FFA500",
    "darkblue": "00008B", "darkred": "8B0000", "darkgreen": "006400",
    "lightgray": "D3D3D3", "lightgrey": "D3D3D3", "gold": "FFD700",
}

_COMMENT_RE = re.compile(r"/\*.*?\*/", re.S)
_RULE_RE = re.compile(r"([^{}]+)\{([^{}]*)\}", re.S)
_ATRULE_RE = re.compile(r"@[a-zA-Z-]+[^{;]*(\{(?:[^{}]|\{[^{}]*\})*\}|;)", re.S)
_LEN_RE = re.compile(r"^\s*(-?[\d.]+)\s*(px|pt|in|cm|mm|pc|em|ex|%)?\s*$", re.I)


def parse_decls(text: str) -> dict:
    """Turn 'color:red; font-size:12pt' into a dict."""
    out = {}
    for chunk in text.split(";"):
        if ":" not in chunk:
            continue
        name, _, value = chunk.partition(":")
        name = name.strip().lower()
        value = value.strip()
        if name in ("padding", "margin") and value:
            # 1-4 values: top, right, bottom, left as CSS repeats them
            v = value.split()
            v = (v * 4)[:4] if len(v) == 1 else (v + v)[:4] if len(v) == 2 else (
                v + [v[1]] if len(v) == 3 else v[:4])
            for side, val in zip(("top", "right", "bottom", "left"), v):
                out[f"{name}-{side}"] = val
        elif name and value:
            out[name] = value
    return out


class StyleSheet:
    """Minimal cascade: tag (1) < class (10) < id (100) < inline (1000)."""

    def __init__(self):
        self.rules = []  # (specificity, order, selector_parts, decls)
        self._order = 0

    def add_css(self, css: str) -> None:
        css = _COMMENT_RE.sub(" ", css)
        css = _ATRULE_RE.sub(lambda m: m.group(1) if m.group(1).startswith("{") else " ", css)
        for selectors, body in _RULE_RE.findall(css):
            decls = parse_decls(body)
            if not decls:
                continue
            for sel in selectors.split(","):
                sel = sel.strip().lower()
                if not sel or " " in sel or ">" in sel:
                    # descendant/child selectors: approximate using the rightmost part
                    sel = re.split(r"[\s>+~]+", sel)[-1]
                    if not sel:
                        continue
                parsed = self._parse_simple(sel)
                if parsed is None:
                    continue
                spec, parts = parsed
                self._order += 1
                self.rules.append((spec, self._order, parts, decls))

    @staticmethod
    def _parse_simple(sel: str):
        m = re.match(r"^([a-z0-9]+)?((?:[.#][\w-]+)*)(?::[\w-]+)?$", sel)
        if not m:
            return None
        tag = m.group(1)
        rest = m.group(2) or ""
        classes = re.findall(r"\.([\w-]+)", rest)
        ids = re.findall(r"#([\w-]+)", rest)
        spec = (1 if tag else 0) + 10 * len(classes) + 100 * len(ids)
        return spec, (tag, classes, ids)

    def declarations_for(self, el: Tag) -> dict:
        tag = el.name.lower()
        el_classes = {c.lower() for c in (el.get("class") or [])}
        el_id = (el.get("id") or "").lower()
        matched = []
        for spec, order, parts, decls in self.rules:
            sel_tag, sel_classes, sel_ids = parts
            if sel_tag and sel_tag != tag:
                continue
            if any(c.lower() not in el_classes for c in sel_classes):
                continue
            if any(i.lower() != el_id for i in sel_ids):
                continue
            matched.append((spec, order, decls))
        matched.sort(key=lambda t: (t[0], t[1]))
        out = {}
        for _, _, decls in matched:
            out.update(decls)
        return out


def computed_style(el: Tag, sheet: StyleSheet, inherited: dict) -> dict:
    """Inherited props + matched rules + presentational attrs + inline style."""
    style = {k: v for k, v in inherited.items() if k in INHERITED_PROPS}
    style.update(sheet.declarations_for(el))

    # legacy presentational attributes that Publisher still emits
    if el.name == "font":
        if el.get("color"):
            style["color"] = el["color"]
        if el.get("face"):
            style["font-family"] = el["face"]
        if el.get("size"):
            try:
                style["font-size"] = {1: "8pt", 2: "10pt", 3: "12pt", 4: "14pt",
                                      5: "18pt", 6: "24pt", 7: "36pt"}[int(el["size"])]
            except (ValueError, KeyError):
                pass
    if el.get("align"):
        style["text-align"] = el["align"]
    if el.get("bgcolor"):
        style["background-color"] = el["bgcolor"]
    if el.name in ("b", "strong"):
        style["font-weight"] = "bold"
    if el.name in ("i", "em"):
        style["font-style"] = "italic"
    if el.name == "u":
        style["text-decoration"] = "underline"
    if el.name in ("s", "strike", "del"):
        style["text-decoration"] = "line-through"

    inline = el.get("style")
    if inline:
        style.update(parse_decls(inline))
    if el.name in ("sup", "sub"):
        style["-x-script"] = "super" if el.name == "sup" else "sub"
    elif str(style.get("vertical-align", "")).lower() in ("super", "sub"):
        style["-x-script"] = str(style["vertical-align"]).lower()
    return style


def to_px(value, dpi: float, pct_basis: Optional[float] = None) -> Optional[float]:
    """CSS length -> pixels (the unit everything else is derived from)."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    m = _LEN_RE.match(str(value))
    if not m:
        return None
    num = float(m.group(1))
    unit = (m.group(2) or "px").lower()
    if unit == "px":
        return num
    if unit == "pt":
        return num * dpi / 72.0
    if unit == "in":
        return num * dpi
    if unit == "cm":
        return num * dpi / 2.54
    if unit == "mm":
        return num * dpi / 25.4
    if unit == "pc":
        return num * dpi / 6.0
    if unit in ("em", "ex"):
        return num * 12.0 * dpi / 72.0 * (0.5 if unit == "ex" else 1.0)
    if unit == "%":
        return None if pct_basis is None else num * pct_basis / 100.0
    return None


def font_pt(value, dpi: float) -> Optional[float]:
    px = to_px(value, dpi)
    return None if px is None else px * 72.0 / dpi


def parse_color(value) -> Optional[RGBColor]:
    if not value:
        return None
    v = str(value).strip().lower()
    if v in ("transparent", "none", "inherit", "initial", "auto"):
        return None
    v = v.replace("window", "").strip() or "x"
    m = re.match(r"^#([0-9a-f]{3})$", v)
    if m:
        h = m.group(1)
        return RGBColor.from_string("".join(c * 2 for c in h).upper())
    m = re.match(r"^#([0-9a-f]{6})", v)
    if m:
        return RGBColor.from_string(m.group(1).upper())
    m = re.match(r"^rgba?\(([^)]+)\)$", v)
    if m:
        parts = [p.strip() for p in m.group(1).split(",")]
        try:
            if len(parts) >= 4 and float(parts[3]) == 0:
                return None
            r, g, b = (int(round(float(p.rstrip('%')) * (2.55 if p.endswith('%') else 1)))
                       for p in parts[:3])
            return RGBColor(max(0, min(255, r)), max(0, min(255, g)), max(0, min(255, b)))
        except ValueError:
            return None
    if v in NAMED_COLORS:
        return RGBColor.from_string(NAMED_COLORS[v])
    m = re.match(r"^([0-9a-f]{6})$", v)
    if m:
        return RGBColor.from_string(m.group(1).upper())
    return None


def background_color(style: dict) -> Optional[RGBColor]:
    c = parse_color(style.get("background-color"))
    if c is not None:
        return c
    bg = style.get("background")
    if bg and "url(" not in bg:
        for token in bg.split():
            c = parse_color(token)
            if c is not None:
                return c
    return None


# ----------------------------------------------------------------------------
# Layout model
# ----------------------------------------------------------------------------

BLOCK_TAGS = {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "li", "blockquote",
              "pre", "dd", "dt", "tr", "address", "center"}
SKIP_TAGS = {"script", "style", "head", "meta", "link", "title", "noscript",
             "object", "param", "v:shapetype", "o:shapedefaults"}

ALIGN_MAP = {
    "left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER,
    "middle": PP_ALIGN.CENTER, "justify": PP_ALIGN.JUSTIFY,
}


@dataclass
class Run:
    text: str
    font: Optional[str] = None
    size_pt: Optional[float] = None
    bold: bool = False
    italic: bool = False
    underline: bool = False
    color: Optional[RGBColor] = None
    caps: Optional[str] = None        # "all" | "small"
    spacing_pt: Optional[float] = None    # letter spacing
    script: Optional[str] = None      # "super" | "sub"


@dataclass
class Para:
    runs: list = field(default_factory=list)
    align: Optional[object] = None
    space_after_pt: float = 0.0
    line_spacing: Optional[float] = None
    indent_pt: float = 0.0            # first line, from the left indent; may be negative
    left_indent_pt: float = 0.0
    right_indent_pt: float = 0.0
    bullet: bool = False              # the export wrote Publisher's bullet as text
    keep_lines: bool = False          # lines broken where Publisher breaks them
    tab_stops: list = field(default_factory=list)    # [(pos_pt, PbTabAlignmentType)]
    pub: Optional[dict] = None        # stage 1's record for this paragraph
    pub_placed: bool = False          # pub was found at this paragraph's place
    best_fit: bool = False            # at Publisher's best-fit size, not the export's
    exact_line_pt: Optional[float] = None
    line_multiple: Optional[float] = None   # PowerPoint's, for lines of mixed sizes

    def text(self) -> str:
        return "".join(r.text for r in self.runs)


@dataclass
class Box:
    """One shape destined for the slide, in CSS pixels."""
    kind: str                      # "text" | "image" | "rect"
    x: float
    y: float
    w: float
    h: float
    paras: list = field(default_factory=list)
    src: Optional[str] = None
    fill: Optional[RGBColor] = None
    line: Optional[RGBColor] = None
    line_w_px: float = 0.0
    order: int = 0
    note: str = ""
    anchor: str = "top"            # text: "top" | "middle" | "bottom"
    fixed: bool = False            # at Publisher's own position: never shifted
    geom: str = "rect"             # rect: a SHAPE_GEOMS key
    adj: tuple = ()                # rect: the preset's adjustment values
    dash: Optional[str] = None     # rect: a LINE_DASHES key
    round_cap: bool = False        # rect: outline drawn with round line ends
    z: Optional[int] = None        # z-index of the positioned shape it came from
    columns: int = 1               # text: columns the text flows through
    column_gap: float = 0.0        # text: space between them


# Font files for measuring inline text. ImageFont.truetype finds these in
# C:\Windows\Fonts by file name. Fonts not listed are looked up in the
# registry's list of installed fonts.
FONT_FILES = {
    "arial": "arial.ttf", "verdana": "verdana.ttf", "calibri": "calibri.ttf",
    "cambria": "cambria.ttc", "times new roman": "times.ttf", "georgia": "georgia.ttf",
    "tahoma": "tahoma.ttf", "trebuchet ms": "trebuc.ttf", "segoe ui": "segoeui.ttf",
    "courier new": "cour.ttf", "garamond": "gara.ttf", "century gothic": "gothic.ttf",
}
# Fonts that are not free to install, replaced by an installed font close in
# style and width when the real one is missing. Keys are lower-case.
FONT_SUBSTITUTES = {
    "abadi": "Gill Sans MT", "abadi extra light": "Gill Sans MT",
    "abadi mt": "Gill Sans MT", "abadi mt condensed": "Gill Sans MT Condensed",
    "abadi mt condensed light": "Gill Sans MT Condensed",
    "abadi mt condensed extra bold": "Gill Sans MT Condensed",
    "elephant pro": "Elephant",
    # Chinese fonts: Publisher draws these with SimSun's Latin letters
    "fangsong": "SimSun", "kaiti": "SimSun",
}
MISSING_FONT = "Calibri"
_font_cache: dict = {}
_installed_fonts: Optional[dict] = None


def installed_fonts() -> dict:
    """Lower-case face name -> font file, from the Windows font registry."""
    global _installed_fonts
    if _installed_fonts is None:
        _installed_fonts = {}
        try:
            import winreg
        except ImportError:
            return _installed_fonts
        key_path = r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"
        for hive in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
            try:
                with winreg.OpenKey(hive, key_path) as key:
                    i = 0
                    while True:
                        try:
                            name, value, _ = winreg.EnumValue(key, i)
                        except OSError:
                            break
                        i += 1
                        name = re.sub(r"\s*\((TrueType|OpenType)\)\s*$", "", name, flags=re.I)
                        # a collection lists its faces as "A & B"
                        for face in name.split(" & "):
                            _installed_fonts.setdefault(face.strip().lower(), str(value))
            except OSError:
                pass
    return _installed_fonts


# the weights Windows uses, nearest first, for a regular face that isn't installed
_REGULAR_STAND_INS = ("light", "medium", "semibold", "bold")


def _font_file(name: str, style: str) -> Optional[str]:
    """The font file for a lower-case face name and style ("bold italic",
    or "" for regular). A family with no regular face (Perpetua Titling MT
    has only Light and Bold) is drawn in its nearest weight."""
    fname = ((installed_fonts().get(f"{name} {style}") if style else None)
             or FONT_FILES.get(name) or installed_fonts().get(name))
    if not fname and not style:
        fname = next((installed_fonts()[f"{name} {w}"] for w in _REGULAR_STAND_INS
                      if f"{name} {w}" in installed_fonts()), None)
    return fname


_family_installed: dict = {}


def font_installed(family: str) -> bool:
    """True if a face of this family is installed, or if the font list is
    unavailable (not Windows)."""
    key = family.lower()
    if key not in _family_installed:
        inst = installed_fonts()
        _family_installed[key] = (not inst or key in inst or key in FONT_FILES
                                  or any(k.startswith(key + " ") for k in inst))
    return _family_installed[key]


def font_name(family: Optional[str]) -> Optional[str]:
    """The family to use. A missing font becomes its substitute, else Calibri,
    which is what Publisher draws in place of any font it does not have."""
    if not family or font_installed(family):
        return family
    sub = FONT_SUBSTITUTES.get(family.lower())
    return sub if sub and font_installed(sub) else MISSING_FONT

# Collapses HTML whitespace but keeps non-breaking spaces, which Publisher
# uses to space out text and inline pictures.
_HTML_WS_RE = re.compile(r"[ \t\r\n\f]+")


# Publisher's HTML export writes a tab as a run of non-breaking spaces sized
# to roughly reach the next tab stop, then a space. Several typed spaces come
# out the same way, so a run counts as tabs only when it ends close to a stop.
_NBSP_RUN_RE = re.compile("(\xa0{2,} ?)")
# Any whitespace run holding a non-breaking space, for laying out inline lines.
_WS_RUN_RE = re.compile("([ \xa0]*\xa0[ \xa0]*)")
TAB_STOP_PT = 36.0               # Publisher's default tab stops: every half inch
INLINE_PICTURE_GAP_PT = 2.88     # Publisher's default spacing around a picture
# Publisher puts a line's extra spacing (above single spacing) below the text,
# or above it for exact spacing; PowerPoint puts about this share of it above.
PPT_EXTRA_ABOVE = 0.66
# Superscript and subscript: both programs draw them at 2/3 size; PowerPoint
# raises (or lowers) them by a share of the run's size, Publisher by about this.
SCRIPT_SCALE = 2.0 / 3.0
SUPER_BASELINE = "42000"
SUB_BASELINE = "-25000"
# Where PowerPoint puts the first baseline below the top of a box, per point of
# font size, for a percentage line spacing (measured; the same for every
# font). Away from single spacing it is linear, three quarters of the extra
# going above the line; just below 100% the line keeps the font's own ascent,
# and so it does until half a percent above.
PPT_BASELINE_SLOPE, PPT_BASELINE_BASE = 0.9014, 0.0063
PPT_BASELINE_NEAR_SINGLE = [(0.8, 0.735), (0.825, 0.771), (0.85, 0.792), (0.875, 0.828),
                            (0.9, 0.845), (0.925, 0.881), (0.95, 0.906), (0.975, 0.942),
                            (1.0, 0.967)]
# Publisher keeps the first baseline where single spacing puts it for a larger
# multiple, but takes this share of a smaller one's cut from above it. Gill
# Sans is the exception: any larger multiple drops it by about the first of
# these shares of the font size (not the Ultra Bold), and for Gill Sans MT
# space after a one-line paragraph at single spacing by the second.
PUB_CUT_ABOVE = 0.75
CJK_LINE_SCALE = 1.3
PUB_GILL_SANS_DROP_MULTIPLE = 0.15
PUB_GILL_SANS_DROP = 0.25
# Gill Sans MT at exact spacing in Publisher: (descent, share of the line
# above the baseline), measured; its font tables give neither
PUB_GILL_SANS_EXACT = (0.33, 0.755)


# ----------------------------------------------------------------------------
# Stage 1's text layout: each paragraph's real text and formatting
# ----------------------------------------------------------------------------

OBJECT_CHAR = "\ufffc"           # Publisher's placeholder for an inline picture
_PUB_TOKEN_RE = re.compile(r"(\s+|\ufffc)")
TAB_ALIGN = {0: "l", 1: "ctr", 2: "r", 3: "dec"}       # PbTabAlignmentType


def pub_key(text: str) -> str:
    """Matches a paragraph across the HTML and stage 1's layout: its words,
    with every kind of whitespace removed."""
    return re.sub(r"\s+", "", text)


# A bullet or list number the export wrote as text (Symbol, Wingdings or
# Unicode glyphs, "1." or "a)"); Publisher's own text leaves them out
_BULLET_RE = re.compile("^[ \xa0]*(?:[•·▪■●○◦§Ø"
                        "ü–‐-]"
                        r"|\(?(?:\d{1,3}|[a-zA-Z]|[ivxIVX]{1,5})[.)])[ \xa0\t]+")


_LONE_BULLETS = set("•·▪■●○◦§Øü–‐-")


def _keep_key_chars(runs: list, n: int) -> None:
    """Cuts a paragraph's runs after its first n non-whitespace characters."""
    for i, run in enumerate(runs):
        if n <= 0:
            del runs[i:]
            return
        seen = 0
        for j, ch in enumerate(run.text):
            if not ch.isspace():
                seen += 1
                if seen == n:
                    run.text = run.text[:j + 1]
                    del runs[i + 1:]
                    return
        n -= seen


def _drop_prefix(runs: list, n: int) -> list:
    """Removes the first n characters of a paragraph's runs and returns them
    as runs of their own, styled as they were."""
    removed = []
    for run in runs:
        if n <= 0:
            break
        cut = min(n, len(run.text))
        removed.append(dataclasses.replace(run, text=run.text[:cut]))
        run.text = run.text[cut:]
        n -= cut
    runs[:] = [r for r in runs if r.text] or runs[:1]
    return removed


def load_pub_text(export_dir: str, base: str) -> dict:
    """Stage 1's <base>_text.json: every paragraph, listed by pub_key. A line
    break splits a paragraph in the HTML, so each line is indexed too."""
    path = os.path.join(export_dir, base + "_text.json")
    try:
        with open(path, encoding="utf-8-sig") as fh:
            data = json.load(fh)
    except (OSError, ValueError):
        return {}
    index: dict = {}

    def add(par, page, frame=None, last=False, box=None):
        lines = re.split("[\v\n]", str(par.get("text", "")))
        tops = sorted(set(par.get("lineTops") or []))
        # each line on a row of its own: its row's top places it
        rows = len(lines) > 1 and par.get("top") is not None and len(tops) == len(lines)
        # with each line's start, every wrapped line's top is known too
        starts = par.get("lineStarts") or []
        all_tops = par.get("lineTops") or []
        lefts = par.get("lineLefts") or []
        if len(lefts) != len(starts):
            lefts = []
        widths = par.get("lineWidths") or []
        if len(widths) != len(lefts):
            widths = []
        wrapped = (len(lines) > 1 and not rows and par.get("top") is not None
                   and len(starts) == len(all_tops))
        offset = 0                             # the line's start, in UTF-16 units
        for i, line in enumerate(lines):
            size = len(line.encode("utf-16-le")) // 2
            key = pub_key(line)
            if key:
                rec = {**par, "text": line, "page": page, "frame": frame, "box": box,
                       "firstIndent": par.get("firstIndent", 0) if i == 0 else 0,
                       # the last text the frame shows; the rest is overflow
                       "overflowAfter": last and i + 1 == len(lines)}
                ks = [k for k, s in enumerate(starts) if offset <= s < offset + max(size, 1)]
                if len(lines) > 1:
                    # this line's own line starts, from its start
                    rec["lineStarts"] = [starts[k] - offset for k in ks]
                    rec["lineLefts"] = [lefts[k] for k in ks] if lefts else []
                    rec["lineWidths"] = [widths[k] for k in ks] if widths else []
                    rec["lineJumps"] = [[s - offset, x] for s, x in par.get("lineJumps") or []
                                        if offset <= s < offset + size]
                    # a line break doesn't end the paragraph: its space
                    # before is above the first line, after below the last
                    rec["spaceBefore"] = par.get("spaceBefore", 0) if i == 0 else 0
                    rec["spaceAfter"] = par.get("spaceAfter", 0) if i + 1 == len(lines) else 0
                if wrapped and ks and starts[ks[0]] == offset:
                    end = all_tops[ks[-1] + 1] if ks[-1] + 1 < len(all_tops) else (
                        par["top"] + (par.get("height") or 0.0))
                    first = all_tops[ks[0]]
                    rec.update(top=first, height=end - first,
                               lineTops=[all_tops[k] for k in ks],
                               spaceBefore=par.get("spaceBefore", 0) if i == 0 else 0,
                               spaceAfter=par.get("spaceAfter", 0) if i + 1 == len(lines) else 0)
                elif wrapped:
                    rec["top"] = None
                elif rows:
                    end = tops[i + 1] if i + 1 < len(tops) else (
                        par["top"] + (par.get("height") or 0.0))
                    rec.update(top=tops[i], height=end - tops[i], lineTops=[tops[i]],
                               spaceBefore=par.get("spaceBefore", 0) if i == 0 else 0,
                               spaceAfter=par.get("spaceAfter", 0) if i + 1 == len(lines) else 0)
                elif len(lines) > 1:
                    rec["top"] = None          # a line's own position isn't recorded
                index.setdefault(key, []).append(rec)
            offset += size + 1

    for shape in data.get("shapes", []):
        try:
            page = int(shape.get("page", 0))    # 0: a master page
        except (TypeError, ValueError):
            page = 0
        pars = _pub_fill_bounds(shape, shape.get("paragraphs") or [])
        frame = _pub_text_area(shape, pars)
        if frame and int(shape.get("columns") or 1) > 1:
            pars, frame = _pub_unfold_columns(pars, frame, int(shape["columns"]))
        # an inline text box's paragraphs have no position: they're found by
        # the box they're in
        box = f"{page}/{shape.get('name')}" if shape.get("inline") else None
        # the last text the frame shows; blank paragraphs after it aren't kept
        shown = max((n for n, par in enumerate(pars) if pub_key(str(par.get("text", "")))),
                    default=-1)
        for n, par in enumerate(pars):
            add(par, page, frame, bool(shape.get("overflowing")) and n == shown, box)
        for cell in shape.get("cells") or []:
            for par in cell.get("paragraphs") or []:
                add(par, page)
    return index


def _pub_set_alike(a: dict, b: dict) -> bool:
    """True if two paragraph records are set in one size and line spacing."""
    return all(a.get(name) == b.get(name) for name in ("size", "lineSpacing", "lineRule"))


def _pub_fill_bounds(shape: dict, pars: list) -> list:
    """Publisher gives a paragraph that runs on past its frame, into a linked
    one for instance, no bounds (top, left and height all 0), though it
    measures its lines in the frame. Their bounds come from the lines
    instead: the top is the first line's, the left the leftmost line's, and
    the bottom the next paragraph's top or, for the last, a line below its
    last line plus its space after. One with a single line and nothing below
    takes the line pitch of the frame's paragraphs set alike, or else stays
    unplaced, and so do lines outside the frame. A linked frame reports
    its first paragraph whole, with the lines it has in the frame before and
    then, from stage 1's version 2, its own: it keeps only those, and the
    text from the first of them, which is what the export gives the frame."""
    out = list(pars)
    try:
        fl, ft = float(shape["left"]), float(shape["top"])
        fr, fb = fl + float(shape["width"]), ft + float(shape["height"])
    except (KeyError, TypeError, ValueError):
        return out

    def inside(t, x):
        return ft - 1.0 <= t <= fb and fl - 1.0 <= x <= fr

    for n in range(len(out) - 1, -1, -1):
        par = out[n]
        tops = par.get("lineTops") or []
        lefts = par.get("lineLefts") or []
        if par.get("top") or par.get("left") or par.get("height") or not tops:
            continue
        starts = par.get("lineStarts") or []
        if n == 0 and len(lefts) == len(tops) == len(starts) and not inside(tops[0], lefts[0]):
            j = next((k for k, (t, x) in enumerate(zip(tops, lefts)) if inside(t, x)), None)
            if j is not None:
                # the text from the frame's first line (starts count UTF-16 units)
                text, cut, units = str(par.get("text", "")), 0, 0
                while cut < len(text) and units < starts[j]:
                    units += len(text[cut].encode("utf-16-le")) // 2
                    cut += 1
                par = dict(par, text=text[cut:], spaceBefore=0.0, firstIndent=0.0,
                           lineStarts=[s - starts[j] for s in starts[j:]],
                           lineJumps=[[s - starts[j], x] for s, x in par.get("lineJumps") or []
                                      if s >= starts[j]])
                for name in ("lineTops", "lineLefts", "lineWidths"):
                    values = par.get(name) or []
                    if len(values) == len(starts):
                        par[name] = values[j:]
                tops, lefts = par["lineTops"], par["lineLefts"]
        if not (ft - 1.0 <= tops[0] <= fb and (not lefts or fl - 1.0 <= min(lefts) <= fr)):
            continue
        # a last line of trailing spaces can sit below the frame
        keep = [k for k, t in enumerate(tops) if t <= fb]
        if len(keep) < len(tops):
            par = dict(par, lineTops=[tops[k] for k in keep])
            for name in ("lineStarts", "lineLefts", "lineWidths"):
                values = par.get(name) or []
                if len(values) == len(tops):
                    par[name] = [values[k] for k in keep]
            tops, lefts = par["lineTops"], par.get("lineLefts") or []
        nxt = out[n + 1].get("top") if n + 1 < len(out) else None
        if nxt:
            bottom = float(nxt)
        elif len(tops) >= 2:
            gaps = sorted(b - a for a, b in zip(tops, tops[1:]))
            bottom = tops[-1] + gaps[len(gaps) // 2] + float(par.get("spaceAfter") or 0.0)
        else:
            gaps = sorted(b - a for q in pars if _pub_set_alike(q, par)
                          for a, b in zip(q.get("lineTops") or [], (q.get("lineTops") or [])[1:]) if b > a)
            if not gaps:
                continue
            bottom = tops[-1] + gaps[len(gaps) // 2] + float(par.get("spaceAfter") or 0.0)
        out[n] = dict(par, top=float(tops[0]), height=bottom - float(tops[0]),
                      left=float(min(lefts)) if lefts else None)
    return out


def _pub_unfold_columns(pars: list, frame: tuple, n: int):
    """A frame's text in n columns as if all in its first column, running on
    down: PowerPoint flows it into columns of its own. Each line moves left
    by its column's offset, and the first line in each later column comes a
    line below the last one in the column before, plus the space between
    the paragraphs where one ends there. Each record keeps where Publisher
    put it on the page, as at, to be found by. Returns the records and the
    first column's text area (left, right in pt)."""
    fl, fr = frame
    span = (fr - fl) / n

    def col(x):
        return min(max(int((x - fl + 1.0) // span), 0), n - 1)

    # each column's offset, from its leftmost line
    base: dict = {}
    for par in pars:
        for x in par.get("lineLefts") or []:
            base[col(x)] = min(base.get(col(x), x), x)
    if 0 not in base or len(base) < 2:
        return pars, frame
    pitch = max((base[k] - base[0]) / k for k in base if k)
    offset = {k: base[k] - base[0] if k in base else k * pitch for k in range(n)}
    out = []
    last = None                  # the line before: (its top, its new top, its column)
    step, after = None, 0.0      # the line pitch, and space after, of the paragraph before
    for par in pars:
        tops, lefts = list(par.get("lineTops") or []), list(par.get("lineLefts") or [])
        if not tops or len(lefts) != len(tops) or par.get("top") is None:
            out.append(par)
            continue
        size = abs(float(par.get("size") or 0.0)) or 12.0
        gaps = sorted(b - a for a, b, x, y in zip(tops, tops[1:], lefts, lefts[1:])
                      if col(x) == col(y) and b > a)
        own = gaps[len(gaps) // 2] if gaps else (step or 1.2 * size)
        new = []
        for i, (t, x) in enumerate(zip(tops, lefts)):
            if last is None:
                v = t
            elif col(x) == last[2]:
                v = last[1] + t - last[0]
            elif i == 0:
                v = last[1] + (step or own) + after + float(par.get("spaceBefore") or 0.0)
            else:
                v = last[1] + own
            new.append(v)
            last = (t, v, col(x))
        step, after = own, float(par.get("spaceAfter") or 0.0)
        top = float(par["top"]) + new[0] - tops[0]
        height = float(par.get("height") or 0.0)
        if len({col(x) for x in lefts}) > 1:
            height = new[-1] + own + after - top
        moved = [x - offset[col(x)] for x in lefts]
        out.append(dict(par, top=top, height=height, lineTops=new, lineLefts=moved,
                        left=min(moved), at=(par.get("left"), par["top"]), columns=n))
    return out, (fl, fr - (n - 1) * pitch)


def _pub_cut(rec: dict, n: int) -> dict:
    """A copy of a paragraph's record cut after its first n non-whitespace
    characters, with the lines that start before the cut. The paragraph goes
    on elsewhere, so its space after doesn't follow the cut."""
    text = str(rec.get("text", ""))
    seen, cut = 0, len(text)
    for i, ch in enumerate(text):
        if not ch.isspace():
            seen += 1
            if seen == n:
                cut = i + 1
                break
    end = len(text[:cut].encode("utf-16-le")) // 2
    starts = rec.get("lineStarts") or []
    keep = [k for k, s in enumerate(starts) if s < end]
    out = dict(rec, text=text[:cut], spaceAfter=0.0,
               lineStarts=[starts[k] for k in keep],
               lineJumps=[j for j in rec.get("lineJumps") or [] if j[0] < end])
    for name in ("lineTops", "lineLefts", "lineWidths"):
        values = rec.get(name) or []
        out[name] = [values[k] for k in keep] if len(values) == len(starts) else []
    tops = out.get("lineTops") or []
    if rec.get("top") is not None and len(tops) >= 2 and len(tops) == len(keep):
        gaps = sorted(b - a for a, b in zip(tops, tops[1:]))
        out["height"] = tops[-1] + gaps[len(gaps) // 2] - float(rec["top"])
    return out


def _pub_unseen_tail(rec: dict, n: int, top: float, left: float) -> Optional[dict]:
    """A copy of a paragraph's record from its last n non-whitespace
    characters on, for the part of a paragraph that runs on into a linked
    text box where stage 1 measured none of its lines: they're only those
    in the box it starts in. Its first line is at top, left (pt), the rest a
    line apart (lineUnseen); how many there are is the next paragraph's
    to say. The first line's gap holds the space before as well. None
    unless every line stage 1 has starts before the cut."""
    text = str(rec.get("text", ""))
    seen, cut = 0, None
    for i in range(len(text) - 1, -1, -1):
        if not text[i].isspace():
            seen += 1
            if seen == n:
                cut = i
                break
    starts = rec.get("lineStarts") or []
    tops = rec.get("lineTops") or []
    if cut is None or not starts or len(tops) != len(starts):
        return None
    if len(text[:cut].encode("utf-16-le")) // 2 <= max(starts):
        return None
    before = float(rec.get("spaceBefore") or 0.0)
    gaps = sorted(b - a - (before if i == 0 else 0.0)
                  for i, (a, b) in enumerate(zip(tops, tops[1:]))
                  if b - a - (before if i == 0 else 0.0) > 0)
    pitch = gaps[len(gaps) // 2] if gaps else None
    size = float(rec.get("size") or 0.0)
    multiple = pub_line_multiple(rec)
    if pitch is None and size > 0 and multiple:
        ratio = pub_single_spacing(rec.get("font"))
        pitch = ratio * size * multiple if ratio else None
    if not pitch:
        return None
    return dict(rec, text=text[cut:], spaceBefore=0.0, firstIndent=0.0,
                lineStarts=[0], lineJumps=[], lineTops=[top], lineLefts=[left],
                lineWidths=[], top=top, left=left, at=None,
                height=pitch + float(rec.get("spaceAfter") or 0.0), lineUnseen=pitch)


def _pub_tail(rec: dict, n: int) -> Optional[dict]:
    """A copy of a paragraph's record from its last n non-whitespace
    characters on, with the lines from there, for the part of a paragraph
    that runs on into a linked text box. None unless a line starts there."""
    text = str(rec.get("text", ""))
    seen, cut = 0, None
    for i in range(len(text) - 1, -1, -1):
        if not text[i].isspace():
            seen += 1
            if seen == n:
                cut = i
                break
    starts = rec.get("lineStarts") or []
    tops, lefts = rec.get("lineTops") or [], rec.get("lineLefts") or []
    if cut is None or len(tops) != len(starts) or len(lefts) != len(starts):
        return None
    begin = len(text[:cut].encode("utf-16-le")) // 2
    if begin not in starts[1:]:
        return None
    k0 = starts.index(begin)
    out = dict(rec, text=text[cut:], spaceBefore=0.0, firstIndent=0.0,
               lineStarts=[s - begin for s in starts[k0:]],
               lineJumps=[[j[0] - begin, j[1]] for j in rec.get("lineJumps") or []
                          if j[0] >= begin],
               lineTops=tops[k0:], lineLefts=lefts[k0:], top=tops[k0], left=lefts[k0],
               at=None)
    widths = rec.get("lineWidths") or []
    out["lineWidths"] = widths[k0:] if len(widths) == len(starts) else []
    gaps = sorted(b - a for a, b in zip(out["lineTops"], out["lineTops"][1:]))
    out["height"] = out["lineTops"][-1] + (gaps[len(gaps) // 2] if gaps else
                                          float(rec.get("size") or 12.0) * 1.2) - tops[k0]
    return out


def _pub_text_area(shape: dict, pars: list):
    """A frame's text area (left, right in pt), where its lines start unless
    they wrap around something; None if unknown. An autoshape keeps its text
    inside its geometry's own inset, which stage 1 doesn't give: there the
    least indented left-aligned line is taken to start at the edge."""
    try:
        margins = shape.get("margins") or [0.0, 0.0, 0.0, 0.0]
        fl = float(shape["left"]) + float(margins[0])
        fr = float(shape["left"]) + float(shape["width"]) - float(margins[2])
    except (KeyError, TypeError, ValueError, IndexError):
        return None
    if shape.get("type", PB_TEXT_FRAME) == PB_TEXT_FRAME:
        return fl, fr
    lead = [x - fl - max(float(par.get("leftIndent") or 0.0), 0.0)
            for par in pars if par.get("align") not in (1, 2)
            for x in (par.get("lineLefts") or [])[1:]]
    if not lead:
        return None
    inset = max(min(lead), 0.0)
    if inset > 0.15 * (fr - fl):
        # too far in for a shape's inset: every line wraps around something
        inset = 0.0
    return fl + inset, fr - inset


def load_pub_frames(export_dir: str, base: str) -> list:
    """Stage 1's text frames that show no text: (page, left, top, width,
    height) in pt. Publisher hides all of such a frame's text as overflow,
    but the export still writes it into the HTML."""
    path = os.path.join(export_dir, base + "_text.json")
    try:
        with open(path, encoding="utf-8-sig") as fh:
            data = json.load(fh)
    except (OSError, ValueError):
        return []
    frames = []
    for shape in data.get("shapes", []):
        if "paragraphs" not in shape or shape.get("cells"):
            continue
        if any(str(par.get("text", "")).strip() for par in shape["paragraphs"]):
            continue
        try:
            frames.append((int(shape.get("page", 0)), float(shape["left"]), float(shape["top"]),
                           float(shape["width"]), float(shape["height"])))
        except (KeyError, TypeError, ValueError):
            continue
    return frames


def load_pub_tables(export_dir: str, base: str) -> list:
    """Stage 1's tables: (page, left, top, width, height) in pt."""
    path = os.path.join(export_dir, base + "_text.json")
    try:
        with open(path, encoding="utf-8-sig") as fh:
            data = json.load(fh)
    except (OSError, ValueError):
        return []
    tables = []
    for shape in data.get("shapes", []):
        if not shape.get("cells"):
            continue
        try:
            tables.append((int(shape.get("page", 0)), float(shape["left"]), float(shape["top"]),
                           float(shape["width"]), float(shape["height"])))
        except (KeyError, TypeError, ValueError):
            continue
    return tables


def load_pub_inline(export_dir: str, base: str) -> Optional[list]:
    """Stage 1's inline text boxes: where Publisher puts each one, in pt, with
    its text's pub_key (inline pictures left out) to match the HTML's copy,
    and whether its text overflows it. None if stage 1 didn't record them
    (layout files before version 2)."""
    path = os.path.join(export_dir, base + "_text.json")
    try:
        with open(path, encoding="utf-8-sig") as fh:
            data = json.load(fh)
    except (OSError, ValueError):
        return None
    if not isinstance(data.get("version"), int) or data["version"] < 2:
        return None
    boxes = []
    for shape in data.get("shapes", []):
        if not shape.get("inline"):
            continue
        text = "".join(str(par.get("text", "")) for par in shape.get("paragraphs") or [])
        try:
            boxes.append({"page": int(shape.get("page", 0)),
                          "box": f"{int(shape.get('page', 0))}/{shape.get('name')}",
                          "key": pub_key(text).replace(OBJECT_CHAR, ""),
                          "left": float(shape["left"]), "top": float(shape["top"]),
                          "width": float(shape["width"]), "height": float(shape["height"]),
                          "overflowing": bool(shape.get("overflowing"))})
        except (KeyError, TypeError, ValueError):
            continue
    return boxes


def _pub_shape_text(shape: dict) -> str:
    pars = list(shape.get("paragraphs") or [])
    for cell in shape.get("cells") or []:
        pars += cell.get("paragraphs") or []
    text = "".join(str(par.get("text", "")) for par in pars)
    return " ".join(text.replace(OBJECT_CHAR, " ").split())


def _pub_text_lines(shape: dict) -> list:
    """The (left, top, right, bottom) of each line of text stage 1 measured
    in a shape, in pt, leaving out blank paragraphs."""
    out = []
    for par in shape.get("paragraphs") or []:
        if not str(par.get("text", "")).replace(OBJECT_CHAR, "").strip():
            continue
        tops, lefts = par.get("lineTops") or [], par.get("lineLefts") or []
        widths = par.get("lineWidths") or []
        try:
            end = float(par["top"]) + float(par["height"]) - float(par.get("spaceAfter") or 0)
        except (KeyError, TypeError, ValueError):
            continue
        for i, top in enumerate(tops):
            if i >= len(lefts) or i >= len(widths) or not widths[i]:
                continue
            bottom = tops[i + 1] if i + 1 < len(tops) else end
            out.append((lefts[i], top, lefts[i] + widths[i], bottom))
    return out


def check_pub_inline(export_dir: str, base: str, page_in: Optional[tuple], warn):
    """Warns about inline text boxes that run off the page or lie on another
    shape's text. Publisher reports where an inline box is but can draw its
    text elsewhere, often over the text around it, so such a box is usually
    leftover text in the source to clean up."""
    path = os.path.join(export_dir, base + "_text.json")
    try:
        with open(path, encoding="utf-8-sig") as fh:
            data = json.load(fh)
    except (OSError, ValueError):
        return
    shapes = [s for s in data.get("shapes", [])
              if all(isinstance(s.get(k), (int, float)) for k in ("left", "top", "width", "height"))]
    parents = {(s.get("page"), s.get("name")): s.get("parent") for s in shapes}
    for shape in shapes:
        text = _pub_shape_text(shape) if shape.get("inline") else ""
        if not text:
            continue
        page = shape.get("page")
        l, t = shape["left"], shape["top"]
        r, b = l + shape["width"], t + shape["height"]
        lines = _pub_text_lines(shape)
        problems = []
        if page_in and any(x1 > page_in[0] * 72 + 6 or y1 > page_in[1] * 72 + 6
                           or x0 < -6 or y0 < -6 for x0, y0, x1, y1 in lines or [(l, t, r, b)]):
            problems.append("runs off the page")
        # an unmeasured box's text runs across the whole box but may stop
        # anywhere down it: only its first line is sure to be there
        own = lines or [(l, t, r, min(b, t + 14))]
        ancestors, up = set(), shape.get("parent")
        while up and up not in ancestors:
            ancestors.add(up)
            up = parents.get((page, up))
        for other in shapes:
            if other is shape or other.get("page") != page or other.get("name") in ancestors \
                    or other.get("parent") == shape.get("name"):
                continue
            if any(min(r1, x1) - max(l1, x0) > 3 and min(b1, y1) - max(t1, y0) > 3
                   for x0, y0, x1, y1 in _pub_text_lines(other) for l1, t1, r1, b1 in own):
                problems.append(f"lies on the text of {other.get('name')!r} "
                                f"({_pub_shape_text(other)[:30]!r})")
        if problems:
            warn(f"page {page}: inline text box {shape.get('name')!r} ({text[:40]!r}) "
                 f"{' and '.join(problems)}; Publisher may draw its text somewhere else "
                 f"or over other text, so check it in the source")


def pub_words_and_gaps(text: str):
    """A paragraph's words (pictures count as words) and the whitespace
    before each word, keyed by word index; the key len(words) is trailing."""
    words, gaps = [], {}
    for tok in _PUB_TOKEN_RE.split(text):
        if not tok:
            continue
        if tok.isspace():
            gaps[len(words)] = gaps.get(len(words), "") + tok
        else:
            words.append(tok)
    return words, gaps


def apply_pub_whitespace(runs: list, pub_text: str) -> bool:
    """Swaps the HTML's whitespace for Publisher's real whitespace, tabs
    included, keeping each run's formatting. False if the words differ."""
    html = "".join(r.text for r in runs)
    h_words, _ = pub_words_and_gaps(html)
    p_words, p_gaps = pub_words_and_gaps(pub_text)
    if h_words != p_words:
        return False
    spans, pos, wi = [], 0, 0                  # (start, end, replacement)
    for tok in _PUB_TOKEN_RE.split(html):
        if not tok:
            continue
        if tok.isspace():
            # Publisher has no space before its first word or after its last
            edge = wi == 0 or wi == len(p_words)
            spans.append((pos, pos + len(tok), p_gaps.get(wi, "" if edge else tok)))
        else:
            wi += 1
        pos += len(tok)
    si, cpos = 0, 0
    for run in runs:
        out = []
        for k, ch in enumerate(run.text):
            g = cpos + k
            while si < len(spans) and spans[si][1] <= g:
                si += 1
            if si < len(spans) and spans[si][0] <= g:
                if g == spans[si][0]:
                    out.append(spans[si][2])
                continue
            out.append(ch)
        cpos += len(run.text)
        run.text = "".join(out)
    return True


_SOFT_HYPHENS = "\x1f\xad"


def _utf16_index(text: str, offset: int) -> int:
    """The str index of a UTF-16 offset, as Publisher counts characters."""
    n = 0
    for i, ch in enumerate(text):
        if n >= offset:
            return i
        n += 2 if ord(ch) > 0xFFFF else 1
    return len(text)


def _pub_line_breaks(para, rec: dict) -> bool:
    """Breaks the paragraph's lines where Publisher does, from stage 1's
    lineStarts, so text wraps the same as in Publisher even where it
    hyphenates a word or wraps around a picture. A word Publisher hyphenates
    gets a hyphen. True if the paragraph is now on Publisher's lines (a
    one-line paragraph already is); False, changing nothing, if the text isn't
    Publisher's or stage 1 didn't record its lines."""
    text = str(rec.get("text", ""))
    starts = rec.get("lineStarts") or []
    full = "".join(r.text for r in para.runs)
    if not starts or full.rstrip() != text.rstrip():
        return False
    breaks = {}                                # str index -> (cut from, insert)
    for s in starts:
        i = _utf16_index(text, s)
        if not 0 < i < len(full.rstrip()):
            continue
        j = i
        # the spaces and tabs at the wrap go: they show nothing, and a tab
        # past the right edge would make PowerPoint wrap again
        while j > 0 and full[j - 1] in " \xa0\t":
            j -= 1
        if j == 0 or full[j - 1] in "\v\n":
            continue
        if j < i:
            breaks[j] = (i, "\v")
        elif full[j - 1] in _SOFT_HYPHENS:
            breaks[j - 1] = (i, "-\v")
        elif full[j - 1].isalnum() and full[i].isalnum():
            breaks[i] = (i, "-\v")            # Publisher hyphenated the word
        else:
            breaks[i] = (i, "\v")
    # a word Publisher moves on past a picture in the middle of its line
    # goes there by a tab, to a stop of its own
    frame = rec.get("frame")
    for s, x in (rec.get("lineJumps") or []) if frame else []:
        i = _utf16_index(text, int(s))
        j = i
        while j > 0 and full[j - 1] in " \xa0":
            j -= 1
        # (after a hyphen, with no space to replace)
        if not 0 < j <= i < len(full.rstrip()) or j in breaks or i in breaks:
            continue
        breaks[j] = (i, "\t")
        para.tab_stops.append((float(x) - frame[0], 0))
    if not breaks:
        return True
    pos, skip_to = 0, 0
    for run in para.runs:
        out = []
        for k, ch in enumerate(run.text):
            g = pos + k
            if g in breaks:
                out.append(breaks[g][1])
                skip_to = breaks[g][0]
                if skip_to > g:
                    continue
            if g < skip_to:
                continue
            out.append(ch)
        pos += len(run.text)
        run.text = "".join(out)
    return True


# Lines starting this far (pt) from where their frame and indents put them
# have wrapped around a picture or another frame
WRAP_SHIFT_PT = 6.0
SMALL_CAPS_SCALE = 0.75                # a small capital's size, of the text's (measured)
PB_TEXT_FRAME = 17                     # PbShapeType of a plain text box


def _split_lines(runs: list) -> list:
    """A paragraph's runs, line by line at its line breaks."""
    lines = [[]]
    for run in runs:
        for k, piece in enumerate(run.text.split("\v")):
            if k:
                lines.append([])
            if piece:
                lines[-1].append(dataclasses.replace(run, text=piece))
    return lines


# The HTML export writes best-fit text at this fraction of the size
# Publisher fits it to
BEST_FIT_EXPORT = 0.95


def _pub_best_fit(p: "Para", rec: dict):
    """Publisher's size for best-fit text, which the HTML export writes 5%
    smaller. Only a paragraph all one size (rec's size is negative when it
    mixes sizes) and exactly 95% of Publisher's in the HTML is changed."""
    try:
        size = float(rec.get("size") or 0.0)
    except (TypeError, ValueError):
        return
    sizes = {r.size_pt for r in p.runs if r.size_pt and r.text.strip()}
    if size <= 0 or len(sizes) != 1:
        return
    html = sizes.pop()
    if abs(html / size - BEST_FIT_EXPORT) > 0.0015:
        return
    for run in p.runs:
        if run.size_pt:
            run.size_pt = round(run.size_pt * size / html, 2)
    p.best_fit = True


def _pub_tracking(p: "Para", rec: dict):
    """Publisher's tracking, which the HTML export leaves out, as letter
    spacing on the runs that have none. Tracking is a percentage of the
    font's average character width (OS/2's), so a narrow font gets less
    space than the font size alone would give it; without the font file,
    each point over 100 adds a 250th of the size."""
    try:
        tracking = float(rec.get("tracking") or 100.0)
    except (TypeError, ValueError):
        return
    # a paragraph with mixed tracking reports nonsense
    if not 0.0 < tracking <= 1000.0 or abs(tracking - 100.0) < 0.5:
        return
    for run in p.runs:
        size = run.size_pt or float(rec.get("size") or 0.0)
        if run.spacing_pt is not None or size <= 0:
            continue
        m = _font_metrics(run.font or rec.get("font"), bool(run.bold), bool(run.italic))
        avg = m["avg_width"] if m and m["avg_width"] > 0 else 0.4
        run.spacing_pt = (tracking - 100.0) / 100.0 * avg * size


def _run_text_width(r, piece: str) -> float:
    """The width of some of a run's text in its font and capitals, without
    its letter spacing."""
    size = r.size_pt or 12.0
    if r.caps == "small":
        # lower case as capitals at a smaller size
        small = "".join(ch for ch in piece if ch.islower())
        return (text_width_px("".join(ch for ch in piece if not ch.islower()), r.font,
                              size, bool(r.bold), bool(r.italic))
                + text_width_px(small.upper(), r.font, size * SMALL_CAPS_SCALE,
                                bool(r.bold), bool(r.italic)))
    return text_width_px(piece.upper() if r.caps else piece, r.font, size,
                         bool(r.bold), bool(r.italic))


def _line_width_pt(line: list) -> float:
    """An estimate of a line's width, without the spaces at its end."""
    width = 0.0
    for r in line:
        piece = r.text.rstrip() if r is line[-1] else r.text
        width += _run_text_width(r, piece) + (r.spacing_pt or 0.0) * len(piece)
    return width


def _pub_wrap_indents(para) -> list:
    """Moves the lines Publisher wraps around a picture or another frame to
    where stage 1 saw them start (lineLefts). PowerPoint has no indent for a
    single line, so a paragraph whose lines start in different places is
    split into one paragraph per run of lines starting alike, on the same
    line spacing. A centred or right-aligned line keeps its alignment within
    the narrower space. Returns the paragraph(s)."""
    rec = para.pub or {}
    frame, lefts = rec.get("frame"), rec.get("lineLefts") or []
    if not para.keep_lines or not frame or not lefts:
        return [para]
    lines = _split_lines(para.runs)
    if len(lines) != len(lefts) or not all(lines):
        return [para]
    fl, fr = frame
    align = rec.get("align")
    widths = rec.get("lineWidths") or []
    if len(widths) != len(lefts):
        widths = []                            # older layout files: estimate them
    pub_text = str(rec.get("text", ""))
    cuts = [_utf16_index(pub_text, s) for s in rec.get("lineStarts") or []] + [len(pub_text)]
    shifts = []                                # (left, right) indent to add, pt
    for j, (line, x) in enumerate(zip(lines, lefts)):
        lead = para.left_indent_pt + (para.indent_pt if j == 0 else 0.0)
        text = "".join(r.text for r in line)
        if align in (1, 2):
            if "\t" in text:
                shifts.append((0.0, 0.0))
                continue
            pub_line = pub_text[cuts[j]:cuts[j + 1]] if j + 1 < len(cuts) else ""
            # a line's bounds run over a run of spaces at its end, out to
            # the frame's edge, though Publisher aligns only what shows
            est = _line_width_pt(line)
            # only an estimate, so only a clear shift counts
            est_tol = max(WRAP_SHIFT_PT, 0.05 * est)
            if widths and len(pub_line) - len(pub_line.rstrip()) < 2:
                width = widths[j]
                tol = WRAP_SHIFT_PT
            else:
                width, tol = est, est_tol
            area_l, area_r = fl + lead, fr - float(rec.get("rightIndent") or 0.0)
            if align == 1:
                s = x + width / 2 - (area_l + area_r) / 2
                # the last line of a big heading can measure half its size
                # too wide on the right, though Publisher centres what shows:
                # a line that starts where a centred one would is not shifted
                if abs(s) > tol and abs(x + est / 2 - (area_l + area_r) / 2) <= est_tol:
                    s = 0.0
                shifts.append((2 * s, 0.0) if s > tol else (0.0, -2 * s) if s < -tol
                              else (0.0, 0.0))
            else:
                gap = area_r - (x + width)
                shifts.append((0.0, gap) if gap > tol else (0.0, 0.0))
            continue
        if j == 0 and para.bullet:
            # stage 1's line starts after Publisher's own bullet
            if para.indent_pt >= 0:
                shifts.append((0.0, 0.0))
                continue
            lead = para.left_indent_pt
        d = x - (fl + lead)
        shifts.append((d, 0.0) if d > WRAP_SHIFT_PT else (0.0, 0.0))
    if not any(l or r for l, r in shifts):
        return [para]
    groups = []                                # [first line, last line, shift]
    for j, sh in enumerate(shifts):
        if groups and all(abs(a - b) < 1.0 for a, b in zip(groups[-1][2], sh)):
            groups[-1][1] = j
        else:
            groups.append([j, j, sh])
    out = []
    for gi, (a, b, (dl, dr)) in enumerate(groups):
        runs = []
        for j in range(a, b + 1):
            line = [dataclasses.replace(r) for r in lines[j]]
            if j > a:
                line[0].text = "\v" + line[0].text
            runs.extend(line)
        out.append(dataclasses.replace(
            para, runs=runs,
            left_indent_pt=para.left_indent_pt + dl,
            indent_pt=para.indent_pt if gi == 0 else 0.0,
            right_indent_pt=para.right_indent_pt + dr,
            space_after_pt=para.space_after_pt if gi + 1 == len(groups) else 0.0))
    return out


# PowerPoint can set small text up to about 2% wider than Publisher (it varies
# with the font and size: hinting rounds each letter's width), so a line
# Publisher just fits would wrap again before its break. Lines measuring within
# this share of the width are tightened to fit, by at most KEEP_LINES_MAX_SPC
# of the font size per character. From KEEP_LINES_LARGE_PT up PowerPoint wraps
# within half a percent of the measured width, and the full share would squeeze
# a large heading visibly narrower than Publisher sets it.
KEEP_LINES_SLACK = 0.025
KEEP_LINES_SLACK_LARGE = 0.01
KEEP_LINES_LARGE_PT = 20.0
KEEP_LINES_MAX_SPC = 0.05
# hinting rounds each letter by a fraction of a pixel, whatever its size: a
# huge heading's share of the width would squeeze it far past what that takes
KEEP_LINES_SLACK_CHAR_PT = 0.3


def _keep_lines_slack(size_pt: float, font: Optional[str] = None) -> float:
    # a fixed-pitch font's letters keep its one width, as Publisher sets them
    m = _font_metrics(font) if font else None
    if m and m.get("fixed"):
        return 0.0
    return KEEP_LINES_SLACK if size_pt < KEEP_LINES_LARGE_PT else KEEP_LINES_SLACK_LARGE


def _keep_lines_extra(width_pt: float, chars: int, size_pt: float, font: Optional[str] = None) -> float:
    """The width (pt) PowerPoint may set a line of text wider than measured."""
    return min(width_pt * _keep_lines_slack(size_pt, font), KEEP_LINES_SLACK_CHAR_PT * chars)

_SPACE_RUN_RE = re.compile(" {8,}")


def _space_breaks(para, rec: dict) -> None:
    """Some authors end lines with a run of spaces long enough to wrap.
    Publisher drops the spaces at the wrap; PowerPoint carries them onto the
    next line. When the runs split the text into exactly Publisher's lines,
    they become line breaks."""
    tops = set(rec.get("lineTops") or [])
    pieces = _SPACE_RUN_RE.split(str(rec.get("text", "")).rstrip())
    if len(tops) < 2 or len(pieces) != len(tops):
        return
    for run in para.runs:
        run.text = _SPACE_RUN_RE.sub("\v", run.text)
    for run in reversed(para.runs):
        run.text = run.text.rstrip(" \v")
        if run.text:
            break


def next_tab_px(x: float, stops_px: list, default_px: float) -> float:
    """Where a tab at x lands: the next custom stop, else the next default one."""
    for stop in stops_px:
        if stop > x + 0.01:
            return stop
    return (int(x // default_px) + 1) * default_px


def pub_line_multiple(rec: dict) -> Optional[float]:
    """A paragraph's line spacing as a multiple, from PbLineSpacingRule."""
    rule = rec.get("lineRule")
    if rule == 0:
        return 1.0
    if rule == 1:
        return 1.5
    if rule == 2:
        return 2.0
    if rule == 5:
        return float(rec.get("lineSpacing") or 1.0)
    return None                  # exact or at-least spacing: keep the HTML's


def ppt_first_baseline(pct: float) -> float:
    """PowerPoint's first baseline below the box top, per point of size."""
    table = PPT_BASELINE_NEAR_SINGLE
    if pct <= table[0][0] or pct > table[-1][0] + 0.005:
        return PPT_BASELINE_BASE + PPT_BASELINE_SLOPE * pct
    for (x0, y0), (x1, y1) in zip(table, table[1:]):
        if pct <= x1:
            return y0 + (y1 - y0) * (pct - x0) / (x1 - x0)
    return table[-1][1]


def pct_lift(line: float, size: float, multiple: float, font: Optional[str],
             bold: bool = False, italic: bool = False,
             pct: Optional[float] = None, measured: bool = False,
             mixed: bool = False, after: float = 0.0) -> Optional[float]:
    """How far to raise a box above Publisher's first line top so that its
    first line, at PowerPoint's percentage spacing line / (1.2 size) (or pct),
    sits where Publisher's does at this multiple. Measured from single
    spacing: line / multiple, unless that is far from the font's single
    spacing and line isn't a measured pitch of text of one size; then the
    font's, any further height being below the text. For a larger multiple
    of the font's own single spacing, from where Publisher puts the first
    baseline at single spacing, which it keeps, where PowerPoint's spacing
    is above 100% and its baseline the same for every font."""
    if size <= 0:
        return None
    pct = pct or line / (1.2 * size)
    ratio = pub_single_spacing(font, bold, italic)
    seen = 1.2 * pct / multiple
    if not ratio or abs(seen / ratio - 1.0) <= 0.03 or (measured and not mixed):
        # the line is Publisher's own single spacing times the multiple (a
        # little off the font's is just rounding), unless text of another
        # size has made it much taller, or it's a lone line's whole height
        ratio = seen
    if not ratio:
        return None
    single = pub_single_spacing(font, bold, italic)
    first = pub_first_baseline(font, bold, italic)
    m = _font_metrics(font, bold, italic)
    if multiple > 1.0 and pct > 1.0 and single and first is not None and abs(ratio / single - 1.0) <= 0.03:
        lift = (ppt_first_baseline(pct) - first) * size
    else:
        lift = (ppt_first_baseline(pct) - ppt_first_baseline(ratio / 1.2)) * size
        if multiple < 1.0:
            lift += PUB_CUT_ABOVE * (1.0 - multiple) * ratio * size
        # where the two put it at single spacing, where that isn't the same
        # place: a font with a line gap. Publisher's line a little off the
        # font's single spacing keeps its baseline, but PowerPoint's moves,
        # and jumps up just past 100%
        gap = (m["typo"] if m["version"] >= 2 else m["hhea"])[2] if m else 0.0
        if single and first is not None and gap > 0:
            ppt = ppt_first_baseline((ratio if abs(ratio / single - 1.0) <= 0.03 else single) / 1.2)
            if single < 1.2:
                ppt = min(ppt, m["win"][0])
            lift -= (first - ppt) * size
    name = (font_name(font) or "").lower()
    if multiple > 1.0 and name.startswith("gill sans") and "ultra" not in name:
        lift -= PUB_GILL_SANS_DROP_MULTIPLE * size
    elif multiple == 1.0 and after > 0 and name == "gill sans mt":
        lift -= PUB_GILL_SANS_DROP * size
    return lift


def snap_to_tab(pos: float, width: float, nbsp_w: float, stop: float) -> Optional[int]:
    """Tabs a whitespace run at pos stands for, or None if it is just spaces."""
    end = pos + width
    k = round(end / stop)
    if k * stop > pos + 0.5 and abs(k * stop - end) <= 2 * nbsp_w:
        return k - int(pos // stop)
    return None


def text_width_px(text: str, family: Optional[str], size_px: float,
                  bold: bool = False, italic: bool = False) -> float:
    """Width of a run of text, measured with the real font where possible."""
    if not text:
        return 0.0
    if HAVE_PIL:
        name = (font_name(family) or "").lower()
        style = " ".join(s for s, on in (("bold", bold), ("italic", italic)) if on)
        key = (name, style, round(size_px * 4))
        if key not in _font_cache:
            font = None
            fname = _font_file(name, style)
            if fname:
                # (PIL's own search of the fonts folder matches case, and
                # Windows names some font files in capitals)
                path = fname if os.path.isabs(fname) else os.path.join(
                    os.environ.get("WINDIR", r"C:\Windows"), "Fonts", fname)
                try:
                    from PIL import ImageFont
                    font = ImageFont.truetype(path if os.path.exists(path) else fname,
                                              max(key[2], 1))
                except (OSError, ImportError):
                    font = None
            _font_cache[key] = font
        font = _font_cache[key]
        if font is not None:
            return font.getlength(text) / 4.0
    # rough fallback: spaces are narrow, other characters about half an em
    spaces = sum(1 for ch in text if ch in " \xa0")
    return (spaces * 0.28 + (len(text) - spaces) * 0.5) * size_px


_metrics_cache: dict = {}


def _font_metrics(family: Optional[str], bold: bool = False,
                  italic: bool = False) -> Optional[dict]:
    """A font's metrics, per point of size: its average character width
    (avg_width), OS/2 typo ascender,
    descender and line gap (typo), the same from hhea (hhea), Windows ascent
    and descent (win), the OS/2 table's version, its code page bits, and
    whether it is fixed-pitch (fixed).
    None if the font file can't be read."""
    name = (font_name(family) or "").lower()
    style = " ".join(s for s, on in (("bold", bold), ("italic", italic)) if on)
    key = (name, style)
    if key in _metrics_cache:
        return _metrics_cache[key]
    found = None
    fname = _font_file(name, style)
    if fname:
        path = fname if os.path.isabs(fname) else os.path.join(
            os.environ.get("WINDIR", r"C:\Windows"), "Fonts", fname)
        try:
            with open(path, "rb") as fh:
                data = fh.read()
            # the first font of a collection
            off = struct.unpack(">I", data[12:16])[0] if data[:4] == b"ttcf" else 0
            tables = {}
            for i in range(struct.unpack(">H", data[off + 4:off + 6])[0]):
                tag, _, at, _ = struct.unpack(">4sIII", data[off + 12 + 16 * i:off + 28 + 16 * i])
                tables[tag] = at
            upm = struct.unpack(">H", data[tables[b"head"] + 18:tables[b"head"] + 20])[0]
            os2, hhea = tables[b"OS/2"], tables[b"hhea"]
            found = dict(
                version=struct.unpack(">H", data[os2:os2 + 2])[0],
                avg_width=struct.unpack(">h", data[os2 + 2:os2 + 4])[0] / upm,
                typo=tuple(v / upm for v in struct.unpack(">hhh", data[os2 + 68:os2 + 74])),
                win=tuple(v / upm for v in struct.unpack(">HH", data[os2 + 74:os2 + 78])),
                hhea=tuple(v / upm for v in struct.unpack(">hhh", data[hhea + 4:hhea + 10])),
                pages=struct.unpack(">I", data[os2 + 78:os2 + 82])[0],
                # post's isFixedPitch, or a Latin text face's PANOSE proportion
                fixed=(b"post" in tables and struct.unpack(
                    ">I", data[tables[b"post"] + 12:tables[b"post"] + 16])[0] != 0)
                or data[os2 + 32] == 2 and data[os2 + 35] == 9)
        except (OSError, KeyError, struct.error, ZeroDivisionError):
            found = None
    _metrics_cache[key] = found
    return found


def pub_single_spacing(family: Optional[str], bold: bool = False,
                       italic: bool = False) -> Optional[float]:
    """Publisher's single line spacing for a font, as a multiple of its size:
    the line height its OS/2 table gives (ascender, descender and line gap),
    or for a font whose OS/2 table is older than version 2, its hhea table.
    PowerPoint's single spacing is 1.2 times the size for every font. None
    if the font file can't be read."""
    m = _font_metrics(family, bold, italic)
    if not m:
        return None
    asc, desc, gap = m["typo"] if m["version"] >= 2 else m["hhea"]
    ratio = (asc - desc + gap) or None
    # an East Asian font (its code pages include Japanese, Chinese or
    # Korean) gets Office's taller lines
    if ratio and m["pages"] & (0b11111 << 17):
        ratio *= CJK_LINE_SCALE
    return ratio


def pub_first_baseline(family: Optional[str], bold: bool = False,
                       italic: bool = False) -> Optional[float]:
    """Where Publisher puts the first baseline at single spacing, below the
    top of the line, per point of size: the ascender and the whole line gap,
    from the table its single spacing comes from (measured). PowerPoint puts
    it where it puts any font's at that percentage, but below 100% no lower
    than the font's Windows ascent: about the same place for a font with no
    line gap, but higher for one with a large gap. None if the font file
    can't be read, or for an East Asian font, not measured."""
    m = _font_metrics(family, bold, italic)
    if not m or m["pages"] & (0b11111 << 17):
        return None
    asc, _, gap = m["typo"] if m["version"] >= 2 else m["hhea"]
    return asc + max(gap, 0.0)


def exact_lift(line: float, size: float, font: Optional[str],
               bold: bool = False, italic: bool = False) -> Optional[float]:
    """How far to raise a box above Publisher's first line top so that its
    first line, at exact spacing line, sits where Publisher's does.
    PowerPoint spaces it like the matching percentage; Publisher puts the
    baseline the font's descent above the line's foot, but no higher than
    its share of the line by Windows ascent and descent."""
    m = _font_metrics(font, bold, italic)
    if size <= 0 or not m:
        return None
    win_asc, win_desc = m["win"]
    desc, share = -m["hhea"][1], win_asc / ((win_asc + win_desc) or 1.0)
    if (font_name(font) or "").lower().startswith("gill sans"):
        desc, share = PUB_GILL_SANS_EXACT
    pub = max(line - desc * size, share * line)
    return ppt_first_baseline(line / (1.2 * size)) * size - pub


def is_positioned(style: dict) -> bool:
    return style.get("position", "static").lower() in ("absolute", "relative", "fixed")


def has_coords(style: dict) -> bool:
    return style.get("left") is not None or style.get("top") is not None


class Converter:
    def __init__(self, args, warn):
        self.dpi = float(args.dpi)
        self.args = args
        self.warn = warn
        # break lines where Publisher does, unless asked to let text reflow
        self.keep_lines = not getattr(args, "reflow", False)
        self.order = 0
        # space after the last block _emit_node placed
        self.trailing = 0.0
        # stage 1's paragraphs, by pub_key; empty for older exports
        self.pub_text: dict = {}
        self.pub_claimed: set = set()     # ids of records already placed
        # stage 1's text frames showing no text, and the page being walked
        self.pub_blank_frames: list = []
        self.pub_tables: list = []         # stage 1's tables (page, left, top, w, h), pt
        self.hide_text = 0                # inside such a frame: no paragraphs
        self.pub_overflowed = False       # the paragraphs just placed end the frame's text
        self.page = 1
        # inline text boxes rebuilt from the VML, by their picture's x-textbox,
        # and where stage 1 saw Publisher put them
        self.inline_boxes: dict = {}
        self.pub_inline: Optional[list] = None
        self.pub_box = None               # the inline text box being placed

    # -- runs and paragraphs ------------------------------------------------

    def _make_run(self, text: str, style: dict) -> Optional[Run]:
        if not text:
            return None
        fam = style.get("font-family")
        if fam:
            fam = font_name(fam.split(",")[0].strip().strip("'\""))
        weight = str(style.get("font-weight", "")).lower()
        bold = weight in ("bold", "bolder") or (weight.isdigit() and int(weight) >= 600)
        deco = str(style.get("text-decoration", "")).lower()
        return Run(
            text=text,
            font=fam or None,
            size_pt=font_pt(style.get("font-size"), self.dpi),
            bold=bold,
            italic=str(style.get("font-style", "")).lower() in ("italic", "oblique"),
            underline="underline" in deco,
            color=parse_color(style.get("color")),
            caps=("all" if str(style.get("text-transform", "")).lower() == "uppercase"
                  else "small" if "small-caps" in str(style.get("font-variant", "")).lower()
                  else None),
            spacing_pt=(lambda v: v * 72.0 / self.dpi if v else None)(
                to_px(style.get("letter-spacing"), self.dpi)),
            script=style.get("-x-script"),
        )

    def collect_paras(self, nodes, sheet: StyleSheet, inherited: dict,
                      stop_nodes: set, region=None, text_area=None) -> list:
        """Paragraphs for a list of sibling nodes (or a single element's children).
        text_area: the text box's own left and right (px), inside its shape's
        geometry, which stage 1's frame for an autoshape only guesses."""
        if self.hide_text:
            return []
        if isinstance(nodes, Tag):
            nodes = list(nodes.children)
        paras: list = []
        current = Para()
        styled: set = set()        # lines an inner block already styled

        def flush():
            nonlocal current
            # Publisher writes an empty paragraph as <p>&nbsp;</p>
            if current.runs and (current.text().strip() or "\xa0" in current.text()):
                paras.append(current)
            current = Para()

        def emit(node, style):
            nonlocal current
            if isinstance(node, PreformattedString):
                return
            if isinstance(node, NavigableString):
                raw = str(node)
                if str(style.get("white-space", "")).lower().startswith("pre"):
                    text = raw.replace("\r\n", "\n")
                else:
                    text = _HTML_WS_RE.sub(" ", raw)
                if not text or (not text.strip() and not current.runs and "\xa0" not in text):
                    return
                run = self._make_run(text, style)
                if run:
                    current.runs.append(run)
                return
            if not isinstance(node, Tag):
                return
            name = node.name.lower()
            if name in SKIP_TAGS or id(node) in stop_nodes:
                return
            if name == "br":
                if current.runs:
                    paras.append(current)
                    current = Para()
                return
            if name in ("img", "table"):
                return
            child_style = computed_style(node, sheet, style)
            if name in BLOCK_TAGS:
                flush()
                start = len(paras)
                for c in node.children:
                    emit(c, child_style)
                # the lines its <br>s split off share the paragraph's alignment,
                # line spacing and left indent; the first line has its indent
                # and the last its space after
                lines = [p for p in paras[start:] if id(p) not in styled]
                if current.runs:
                    lines.append(current)
                styled.update(id(p) for p in lines)
                align = ALIGN_MAP.get(str(child_style.get("text-align", "")).lower())
                lh = to_px(child_style.get("line-height"), self.dpi)
                fs = to_px(child_style.get("font-size", "12pt"), self.dpi)
                ml = to_px(child_style.get("margin-left"), self.dpi)
                for line in lines:
                    line.align = align
                    if lh and fs:
                        line.line_spacing = round(min(max(lh / fs, 0.5), 3.0), 3)
                    line.left_indent_pt = max(ml or 0.0, 0.0) * 72.0 / self.dpi
                if lines:
                    ti = to_px(child_style.get("text-indent"), self.dpi)
                    lines[0].indent_pt = (ti or 0) * 72.0 / self.dpi
                if current.runs:
                    mb = to_px(child_style.get("margin-bottom"), self.dpi)
                    current.space_after_pt = (mb or 0) * 72.0 / self.dpi
                flush()
            else:
                for c in node.children:
                    emit(c, child_style)

        for n in nodes:
            emit(n, inherited)
        flush()
        base_align = ALIGN_MAP.get(str(inherited.get("text-align", "")).lower())
        for p in paras:
            if p.align is None:
                p.align = base_align
            rec, placed = self.pub_find(pub_key(p.text()), region)
            # the export writes a bullet as text; Publisher's text has none
            bullet = _BULLET_RE.match(p.text()) if rec is None else None
            bullet_runs = []
            if bullet:
                rec, placed = self.pub_find(pub_key(p.text()[bullet.end():]), region)
                if rec:
                    bullet_runs = _drop_prefix(p.runs, bullet.end())
            if not placed and region is not None:
                # a paragraph running into the overflow: Publisher's text is
                # only its visible start, so keep just that (the whole text
                # may be another frame's, elsewhere)
                start, at = self.pub_find_start(pub_key(p.text()), region)
                if start:
                    rec, placed = start, at
                    _keep_key_chars(p.runs, len(pub_key(rec["text"])))
                else:
                    # a paragraph running on into a linked text box
                    head, at = self.pub_find_continued(pub_key(p.text()), region)
                    if not head:
                        # or the rest of one that started in another
                        first = next((q for q in paras if q.text().strip()), None) is p
                        head, at = self.pub_find_tail(pub_key(p.text()), region,
                                                      text_area if first else None)
                    if head:
                        rec, placed = head, at
            if rec and rec.get("frame") and text_area:
                # lines are indented from where the box's text starts
                rec = dict(rec, frame=tuple(v * 72.0 / self.dpi for v in text_area))
            if rec and apply_pub_whitespace(p.runs, rec["text"]):
                p.tab_stops = [(t.get("pos", 0.0), t.get("align", 0))
                               for t in rec.get("tabs") or []]
                p.pub, p.pub_placed = rec, placed
                _pub_best_fit(p, rec)
                p.bullet = bool(bullet_runs)
                # (an unseen tail's line breaks aren't known)
                p.keep_lines = (self.keep_lines and not rec.get("lineUnseen")
                                and _pub_line_breaks(p, rec))
                if not p.keep_lines:
                    _space_breaks(p, rec)
                # Publisher's indents: the HTML leaves out hanging ones
                p.left_indent_pt = max(float(rec.get("leftIndent") or 0.0), 0.0)
                p.indent_pt = max(float(rec.get("firstIndent") or 0.0), -p.left_indent_pt)
                _pub_tracking(p, rec)
            elif p.align in (None, PP_ALIGN.LEFT):
                self._restore_tabs(p)
            if bullet_runs:
                if p.indent_pt < 0:
                    # a hanging bullet: a tab takes the text to the indent
                    tab = bullet_runs[-1]
                    if not tab.text.strip(" \xa0\t") and p.runs:
                        # in the text's font: the export leaves the gap unsized
                        tab.font, tab.size_pt = p.runs[0].font, p.runs[0].size_pt
                    tab.text = tab.text.rstrip(" \xa0\t") + "\t"
                    bullet_runs = [r for r in bullet_runs if r.text]
                p.runs[:0] = bullet_runs
        text_paras = [p for p in paras if p.text().strip()]
        if not text_paras or all(p.pub_placed for p in text_paras):
            # _pub_place folds the empty paragraphs into the space after
            return text_paras
        # otherwise an empty paragraph is a blank line, except at the end
        while paras and not paras[-1].text().strip():
            paras.pop()
        return paras

    def pub_find(self, key: str, region=None):
        """Stage 1's record for a paragraph: the first unclaimed one inside
        region (x0, y0, x1, y1 in px) if there is one, so repeated text matches
        the right copy, and claims it. Returns (record, found_in_region)."""
        recs = self.pub_text.get(key)
        if not recs:
            return None, False
        if self.pub_box is not None:
            for rec in recs:
                if rec.get("box") == self.pub_box and id(rec) not in self.pub_claimed:
                    self.pub_claimed.add(id(rec))
                    # placed if stage 1 could measure the box's lines
                    return rec, rec.get("top") is not None and rec.get("left") is not None
        if region is not None:
            x0, y0, x1, y1 = region
            tol = 3.0
            for rec in recs:
                if rec.get("top") is None or rec.get("left") is None:
                    continue
                if id(rec) in self.pub_claimed or rec.get("page", 0) not in (0, self.page):
                    continue
                rx, ry = (v * self.dpi / 72.0 for v in rec.get("at") or (rec["left"], rec["top"]))
                if not (x0 - tol <= rx <= x1 + tol and y0 - tol <= ry <= y1 + tol):
                    continue
                # a frame lying over a bigger one: its text isn't the bigger
                # one's (which may hold a hidden copy of it in its overflow)
                frame = rec.get("frame")
                if frame and ((frame[1] - frame[0]) * (rec.get("columns") or 1)
                              * self.dpi / 72.0 < 0.6 * (x1 - x0)):
                    continue
                self.pub_claimed.add(id(rec))
                return rec, True
        return recs[0], False

    def pub_blank_frame(self, x: float, y: float, w: float, h: float) -> bool:
        """True if stage 1 saw a text frame here (px) showing no text: all
        the text the HTML gives it is overflow, which Publisher hides."""
        k = self.dpi / 72.0
        tol = 2.0
        for page, left, top, width, height in self.pub_blank_frames:
            if page not in (0, self.page):
                continue
            if (abs(left * k - x) <= tol and abs(top * k - y) <= tol
                    and abs(width * k - w) <= tol and abs(height * k - h) <= tol):
                return True
        return False

    def pub_find_start(self, key: str, region):
        """An unclaimed record inside region whose text is the start of key,
        for a paragraph cut short by its text box's overflow. Claims it. A
        short start only where stage 1 says the frame's text ends."""
        x0, y0, x1, y1 = region
        tol = 3.0
        for k, recs in self.pub_text.items():
            if not k or len(k) >= len(key) or not key.startswith(k):
                continue
            for rec in recs:
                if (rec.get("top") is None or id(rec) in self.pub_claimed
                        or rec.get("page", 0) not in (0, self.page)
                        or (len(k) < 20 and not rec.get("overflowAfter"))):
                    continue
                rx, ry = (v * self.dpi / 72.0 for v in rec.get("at") or (rec["left"], rec["top"]))
                if x0 - tol <= rx <= x1 + tol and y0 - tol <= ry <= y1 + tol:
                    self.pub_claimed.add(id(rec))
                    return rec, True
        return None, False

    def pub_find_continued(self, key: str, region):
        """An unclaimed record inside region whose text starts with key, for a
        paragraph that runs on into a linked text box: the export splits its
        text between the boxes, while stage 1 records all of it in the first,
        with the lines it shows there. Claims it and returns a copy cut to
        key's text."""
        if len(key) < 20:
            return None, False
        x0, y0, x1, y1 = region
        tol = 3.0
        for k, recs in self.pub_text.items():
            if len(k) <= len(key) or not k.startswith(key):
                continue
            for rec in recs:
                if (rec.get("top") is None or rec.get("left") is None
                        or id(rec) in self.pub_claimed
                        or rec.get("page", 0) not in (0, self.page)):
                    continue
                rx, ry = (v * self.dpi / 72.0 for v in rec.get("at") or (rec["left"], rec["top"]))
                if x0 - tol <= rx <= x1 + tol and y0 - tol <= ry <= y1 + tol:
                    self.pub_claimed.add(id(rec))
                    return _pub_cut(rec, len(key)), True
        return None, False

    def pub_find_tail(self, key: str, region, text_area=None):
        """A record whose text ends with key, its lines from there inside
        region, for the part of a paragraph that runs on into a linked text
        box, where the export starts a paragraph of its own. Returns a copy cut
        to key's text; the record isn't claimed (its start is another box's).
        With text_area (the box's first paragraph), one whose lines stage 1
        measured only before key's text, in another box, starts at the top
        of this one's text, as far below its top as text_area is in from
        its left. A short key only there, and only if one paragraph on the
        page ends with it."""
        if len(key) < 20:
            if not text_area or len(key) < 8:
                return None, False
            ends = {k for k, recs in self.pub_text.items()
                    if len(k) > len(key) and k.endswith(key)
                    and any(r.get("page", 0) in (0, self.page) for r in recs)}
            if len(ends) != 1:
                return None, False
        x0, y0, x1, y1 = region
        tol = 3.0
        k72 = 72.0 / self.dpi
        unseen = None
        for k, recs in self.pub_text.items():
            if len(k) <= len(key) or not k.endswith(key):
                continue
            for rec in recs:
                if rec.get("page", 0) not in (0, self.page):
                    continue
                tail = _pub_tail(rec, len(key))
                if tail is None:
                    if text_area and unseen is None:
                        inset = max(text_area[0] - x0, 0.0)
                        unseen = _pub_unseen_tail(rec, len(key), (y0 + inset) * k72,
                                                  text_area[0] * k72)
                        lines = zip(rec.get("lineLefts") or [], rec.get("lineTops") or [])
                        if unseen and any(x0 - tol <= lx / k72 <= x1 + tol and
                                          y0 - tol <= ly / k72 <= y1 + tol for lx, ly in lines):
                            unseen = None   # its lines are in this box after all
                    continue
                rx, ry = tail["left"] * self.dpi / 72.0, tail["top"] * self.dpi / 72.0
                if x0 - tol <= rx <= x1 + tol and y0 - tol <= ry <= y1 + tol:
                    return tail, True
        return (unseen, True) if unseen else (None, False)

    def _pub_place(self, paras: list, bottom: Optional[float] = None):
        """Puts paragraphs where Publisher puts them: sets each one's exact line
        spacing and space after so it starts at Publisher's top. Returns (top,
        height) in px, or None unless every paragraph was found in place.

        Empty paragraphs are left out (their space falls into the space after),
        and so is a text box's overflow, which Publisher hides: paragraphs at
        the end that stage 1 did not see in this box, below text that reaches
        the frame's bottom (px). Both are removed from paras."""
        text = [p for p in paras if p.text().strip()]
        # an empty bulleted paragraph: the export writes its bullet, but
        # stage 1 skips it as empty. It fills the gap between the ones
        # around it.
        for i in range(1, len(text) - 1):
            p, before, after = text[i], text[i - 1].pub, text[i + 1].pub
            if (p.pub_placed or not text[i - 1].pub_placed or not text[i + 1].pub_placed
                    or p.text().strip(" \xa0\t") not in _LONE_BULLETS):
                continue
            top = before["top"] + (before.get("height") or 0.0)
            if after["top"] - top <= 1.0:
                continue
            p.pub = dict(before, text="", top=top, height=after["top"] - top, lineTops=[top],
                         lineStarts=[0], lineLefts=[], lineWidths=[], lineJumps=[],
                         spaceBefore=0.0, spaceAfter=0.0, overflowAfter=False)
            p.pub_placed = True
        if bottom is not None:
            n = len(text)
            while n and not text[n - 1].pub_placed:
                n -= 1
            if 0 < n < len(text) and text[n - 1].pub_placed:
                last = text[n - 1].pub
                lines = max(len(last.get("lineTops") or []), 1)
                end = (last["top"] + (last.get("height") or 0.0)) * self.dpi / 72.0
                # stage 1 says the frame overflows, or the text reaches its bottom
                if last.get("overflowAfter") or \
                        end >= bottom - 2.5 * (last.get("height") or 0.0) / lines * self.dpi / 72.0:
                    self.warn(f"dropped {len(text) - n} paragraph(s) in a text box's "
                              f"overflow, which Publisher hides: {text[n].text()[:40]!r}")
                    # drop them even when the rest can't be placed below
                    del paras[next(i for i, p in enumerate(paras) if p is text[n]):]
                    text = text[:n]
        if not text or not all(p.pub_placed for p in text):
            return None
        # the last text the frame shows: what follows it in the frame is
        # in its overflow too
        self.pub_overflowed = bool(text[-1].pub.get("overflowAfter"))
        paras[:] = text
        for i, p in enumerate(paras):
            pitch = p.pub.get("lineUnseen")
            if not pitch or i + 1 >= len(paras):
                continue
            rec, nxt = p.pub, paras[i + 1].pub
            after = float(rec.get("spaceAfter") or 0.0)
            n = max(int(round((nxt["top"] - (nxt.get("spaceBefore") or 0.0) - after
                               - rec["top"]) / pitch)), 1)
            p.pub = dict(rec, lineTops=[rec["top"] + k * pitch for k in range(n)],
                         height=n * pitch + after)

        def text_tops(p):
            """The paragraph's line tops, and how far its text's top is below
            Publisher's: a first line made tall by something in it, such as
            a picture, has its text on the line's foot, a line above the
            next one. (A bullet makes it tall in PowerPoint too.)"""
            rec = p.pub
            # older layout files can repeat the last line
            tops = sorted(set(rec.get("lineTops") or [])) or [rec["top"]]
            # the first line's top is above the paragraph's space before,
            # which Publisher keeps at the top of a frame or cell too
            before = rec.get("spaceBefore") or 0.0
            if before > 0 and (len(tops) < 2 or tops[1] - tops[0] > before):
                tops = [tops[0] + before] + tops[1:]
            else:
                before = 0.0
            gaps = [b - a for a, b in zip(tops, tops[1:])]
            if len(gaps) >= 2 and not p.bullet:
                rest = sorted(gaps[1:])[len(gaps[1:]) // 2]
                if rest > 0 and gaps[0] > 1.5 * rest:
                    shift = gaps[0] - rest
                    return [tops[0] + shift] + tops[1:], shift + before
            return tops, before

        settings = []
        for i, p in enumerate(paras):
            rec = p.pub
            tops, _ = text_tops(p)
            height = rec.get("height") or 0.0
            # the next paragraph's top; an empty paragraph between them isn't
            # in the HTML, so its space ends up in this one's space after
            nxt = (paras[i + 1].pub["top"] + text_tops(paras[i + 1])[1] if i + 1 < len(paras)
                   else rec["top"] + height)
            if len(tops) >= 2:
                line = (tops[-1] - tops[0]) / (len(tops) - 1)
            else:
                line = height - (rec.get("spaceAfter") or 0.0) - (rec.get("spaceBefore") or 0.0)
            # how far PowerPoint would put this paragraph's text below Publisher's
            multiple = pub_line_multiple(rec) if rec.get("lineRule") in (0, 1, 2, 5) else None
            sized = [r for r in p.runs if r.text.strip() and r.size_pt and not r.script]
            if len(tops) < 2 and multiple and sized:
                # a single line's height comes from the paragraph's, which can
                # leave out some of its space after: too short for its text,
                # take the font's
                big = max(sized, key=lambda r: r.size_pt)
                ratio = pub_single_spacing(big.font, big.bold, big.italic)
                if ratio and line < 0.9 * ratio * big.size_pt * multiple:
                    line = ratio * big.size_pt * multiple
            if line <= 0:
                return None
            p.line_multiple = None
            line_sizes = _line_size_list(p) if p.keep_lines and multiple else []
            # (no gap measures the last line, so its size must be one the
            # others show)
            if (len(tops) >= 2 and len(line_sizes) == len(tops) and None not in line_sizes
                    and len(set(line_sizes[:-1])) > 1 and line_sizes[-1] in line_sizes[:-1]):
                # lines of different sizes, each as tall as its largest text:
                # a multiple in PowerPoint too, which spaces lines the same way
                pitches = [(b - a, 1.2 * s) for a, b, s in zip(tops, tops[1:], line_sizes)]
                fit = sum(g for g, _ in pitches) / sum(s for _, s in pitches)
                if all(abs(g - fit * s) <= 0.6 for g, s in pitches):
                    p.line_multiple = fit
                    # the last line's height, for the space after it
                    line = fit * 1.2 * line_sizes[-1]
            lift = PPT_EXTRA_ABOVE * line * (1.0 - 1.0 / multiple) if multiple and multiple > 1 else 0.0
            sizes = [r.size_pt for r in p.runs if r.size_pt and r.text]
            first_sizes = _line_size_list(p)[:1]
            if multiple and sized and (p.line_multiple or len(_line_sizes(p)) == 1)                     and first_sizes and first_sizes[0]:
                # PowerPoint spaces the lines by a percentage: where that puts
                # the first line, against where Publisher's multiple does
                big = max(sized, key=lambda r: r.size_pt)
                better = pct_lift(line, first_sizes[0], multiple, big.font, big.bold, big.italic,
                                  pct=p.line_multiple, measured=len(tops) >= 2,
                                  mixed=(rec.get("size") or 0) < 0,
                                  after=(rec.get("spaceAfter") or 0.0) if len(tops) < 2 else 0.0)
                if better is not None:
                    lift = better
            if rec.get("lineRule") == 3 and sizes:
                # exact spacing: Publisher puts the extra above the text
                big = max(sized, key=lambda r: r.size_pt) if sized else None
                exact = (exact_lift(line, first_sizes[0], big.font, big.bold, big.italic)
                         if big and first_sizes and first_sizes[0] else None)
                # (rec's size is -9999999 when the paragraph mixes sizes.)
                lift = exact if exact is not None else \
                    -(1.0 - PPT_EXTRA_ABOVE) * (line - 1.2 * max(sizes))
            settings.append([line, nxt - tops[-1] - line, lift])
        # raise each paragraph by its lift: the box by the first one's, and
        # each later one through the space after the paragraph before it
        for i in range(1, len(settings)):
            settings[i - 1][1] -= settings[i][2] - settings[i - 1][2]
        for p, (line, after, _) in zip(paras, settings):
            p.exact_line_pt = line
            p.space_after_pt = max(after, 0.0)
        first, last = paras[0].pub, paras[-1].pub
        first_shift = text_tops(paras[0])[1]
        paras[:] = [q for p in paras for q in _pub_wrap_indents(p)]
        k = self.dpi / 72.0
        top = first["top"] + first_shift - settings[0][2]
        return top * k, (last["top"] + (last.get("height") or 0.0) - top) * k

    def _restore_tabs(self, para: Para) -> None:
        """Turns the export's tab runs back into tabs, measuring from the frame
        edge as Publisher does. Assumes the runs fall on the first line."""
        stop = TAB_STOP_PT * self.dpi / 72.0
        pos = para.indent_pt * self.dpi / 72.0
        for run in para.runs:
            size_px = (run.size_pt or 12.0) * self.dpi / 72.0
            nbsp_w = text_width_px("\xa0", run.font, size_px)
            out = []
            for i, piece in enumerate(_NBSP_RUN_RE.split(run.text)):
                width = text_width_px(piece, run.font, size_px)
                tabs = snap_to_tab(pos, width, nbsp_w, stop) if i % 2 else None
                if tabs:
                    out.append("\t" * tabs)
                    pos = (int(pos // stop) + tabs) * stop
                else:
                    out.append(piece)
                    pos += width
            run.text = "".join(out)

    def estimate_height(self, paras: list, width_px: float) -> float:
        """Rough laid-out height, used only to stack sibling content sensibly."""
        total = 0.0
        width_px = max(width_px, 20.0)
        for p in paras:
            size_px = max((r.size_pt or 12.0) for r in p.runs) * self.dpi / 72.0
            chars = len(p.text())
            per_line = max(int(width_px / max(size_px * 0.5, 1.0)), 1)
            lines = max(1, -(-chars // per_line))
            lead = (p.line_spacing or 1.18)
            total += lines * size_px * lead + p.space_after_pt * self.dpi / 72.0
        return total

    # -- geometry -----------------------------------------------------------

    def _box_geometry(self, el: Tag, style: dict, origin, cb_w, cb_h):
        left = to_px(style.get("left"), self.dpi, cb_w)
        top = to_px(style.get("top"), self.dpi, cb_h)
        w = to_px(style.get("width"), self.dpi, cb_w)
        h = to_px(style.get("height"), self.dpi, cb_h)

        if w is None and el.get("width"):
            w = to_px(el["width"], self.dpi, cb_w)
        if h is None and el.get("height"):
            h = to_px(el["height"], self.dpi, cb_h)

        right = to_px(style.get("right"), self.dpi, cb_w)
        bottom = to_px(style.get("bottom"), self.dpi, cb_h)
        if left is None and right is not None and w is not None:
            left = cb_w - right - w
        if top is None and bottom is not None and h is not None:
            top = cb_h - bottom - h

        return origin[0] + (left or 0.0), origin[1] + (top or 0.0), w, h

    # -- tables -------------------------------------------------------------

    def _table_boxes(self, table: Tag, sheet: StyleSheet, inherited: dict,
                     x, y, w, h, stop_nodes) -> tuple:
        """Returns (boxes, consumed_height). Cells become individual text boxes."""
        rows = [tr for tr in table.find_all("tr") if tr.find(["td", "th"])]
        if not rows:
            return [], 0.0
        # a table in a group: the export's wrapper can be off from where
        # Publisher puts it, which stage 1 recorded
        tw_, th_ = to_px(table.get("width"), self.dpi), to_px(table.get("height"), self.dpi)
        k = self.dpi / 72.0
        for page, left, top, width, height in self.pub_tables:
            if (page == self.page and tw_ and th_ and abs(width * k - tw_) <= 2.0
                    and abs(height * k - th_) <= 2.0 and abs(top * k - y) <= 3.0
                    and 1.0 < abs(left * k - x) <= 0.25 * tw_):
                x, y = left * k, top * k
                break

        pad = to_px(table.get("cellpadding"), self.dpi)
        pad = 2.0 if pad is None else pad
        grid = [r.find_all(["td", "th"]) for r in rows]
        ncols = max(sum(max(1, int(c.get("colspan", 1) or 1)) for c in cells) for cells in grid)
        ncols = max(ncols, 1)
        w = w or to_px(table.get("width"), self.dpi) or 400.0

        # column widths from the first row with a cell per column, since a
        # cell spanning columns doesn't say how its width divides
        col_w = [w / ncols] * ncols
        for cells in grid:
            if len(cells) != ncols:
                continue
            declared = [to_px(computed_style(c, sheet, inherited).get("width") or c.get("width"),
                              self.dpi, w) for c in cells]
            if all(d for d in declared):
                total = sum(declared)
                col_w = [d * w / total for d in declared]
                break

        # row heights from declared values, else from estimated content
        row_h, cell_paras = [], []
        for ri, cells in enumerate(grid):
            rs = computed_style(rows[ri], sheet, inherited)
            declared_h = to_px(rs.get("height") or rows[ri].get("height"), self.dpi)
            per_row, tallest = [], 0.0
            ci = 0
            for cell in cells:
                span = max(1, int(cell.get("colspan", 1) or 1))
                cw = sum(col_w[ci:ci + span]) if ci < ncols else col_w[-1]
                cs = computed_style(cell, sheet, inherited)
                # Publisher sets the height on the cells, not the row
                cell_h = to_px(cs.get("height") or cell.get("height"), self.dpi)
                if cell_h and int(cell.get("rowspan", 1) or 1) == 1:
                    declared_h = max(declared_h or 0.0, cell_h)
                if cell.name == "th":
                    cs.setdefault("font-weight", "bold")
                paras = self.collect_paras(cell, sheet, cs, stop_nodes)
                est = self.estimate_height(paras, cw - 2 * pad) + 2 * pad
                tallest = max(tallest, est)
                per_row.append((cell, cs, paras, ci, cw))
                ci += span
            cell_paras.append(per_row)
            row_h.append(declared_h or max(tallest, 14.0))

        if h and sum(row_h) < h * 0.92:
            pass  # declared container taller than the content: leave rows compact
        boxes = []
        # each cell's borders by edge, drawn once over all the cell fills
        edges: dict = {}
        table_box = (x, y, sum(col_w), sum(row_h))
        cy = y
        for ri, per_row in enumerate(cell_paras):
            for cell, cs, paras, ci, cw in per_row:
                cx = x + sum(col_w[:ci])
                fill = background_color(cs)
                if fill is not None:
                    self.order += 1
                    boxes.append(Box(kind="rect", x=cx, y=cy, w=cw, h=row_h[ri],
                                     fill=fill, order=self.order, note="table cell fill"))
                self._cell_borders(cs, cx, cy, cw, row_h[ri], table_box, edges)
                # HTML centres cell content vertically unless told otherwise,
                # and Publisher's export relies on that
                valign = str(cs.get("vertical-align") or cell.get("valign")
                             or "middle").lower()
                if valign not in ("top", "bottom"):
                    valign = "middle"
                if cell.find(["img", "table"]) is not None:
                    # pictures or a nested table: lay the cell out like a container
                    cell_style = {k: v for k, v in cs.items()
                                  if not k.startswith(("background", "border"))}
                    placed: list = []
                    used = self._emit_node(cell, sheet, cell_style, (cx, cy), cw, row_h[ri],
                                           placed, stop_nodes)
                    pad_b = to_px(cs.get("padding-bottom"), self.dpi) or 0.0
                    # like Publisher, centre without the last paragraph's space after
                    free = max(row_h[ri] - (used - self.trailing) - pad_b, 0.0)
                    dy = {"top": 0.0, "middle": free / 2, "bottom": free}[valign]
                    for b in placed:
                        if not b.fixed:
                            b.y += dy
                    # Publisher hides a text box's overflow: whatever starts
                    # below its bottom, and inline pictures too wide to fit.
                    # One that leads the frame, where stage 1 saw it laid out,
                    # Publisher draws running past the frame's edge; further
                    # down it draws them jumbled over the frame's other lines
                    bottom = cy + row_h[ri]
                    first = min((b.y for b in placed), default=0.0)
                    kept = [b for b in placed
                            if b.y < bottom - 1.0
                            and not (b.kind == "image" and b.w > cw + 2.0
                                     and not (b.fixed and b.y <= first + 1.0))]
                    if len(kept) < len(placed):
                        self.warn(f"dropped {len(placed) - len(kept)} shape(s) in a text box's "
                                  f"overflow, which Publisher hides")
                    boxes.extend(kept)
                elif paras:
                    # the cell's own padding, else the table's cellpadding,
                    # plus that of the text box's <div class=shape> inside it
                    pl, pt_, pr, pb = (
                        (pad if v is None else v) + inner for v, inner in zip(
                            (to_px(cs.get("padding-" + side), self.dpi, cw)
                             for side in ("left", "top", "right", "bottom")),
                            self._wrapper_padding(cell, sheet, cs, cw)))
                    # (each frame of a linked story records all its text, so
                    # a record's own frame may be another one's)
                    paras = self.collect_paras(cell, sheet, cs, stop_nodes,
                                               region=(cx, cy, cx + cw, cy + row_h[ri]),
                                               text_area=(cx + pl, cx + max(cw - pr, pl + 8.0)))
                    where = self._pub_place(paras, bottom=cy + row_h[ri])
                    self.order += 1
                    if where:
                        top, _ = where
                        boxes.append(Box(kind="text", x=cx + pl, y=top,
                                         w=max(cw - pl - pr, 8.0),
                                         h=max(cy + row_h[ri] - top, 8.0),
                                         paras=paras, order=self.order, note="table cell",
                                         fixed=True))
                    else:
                        boxes.append(Box(kind="text", x=cx + pl, y=cy + pt_,
                                         w=max(cw - pl - pr, 8.0),
                                         h=max(row_h[ri] - pt_ - pb, 8.0),
                                         paras=paras, order=self.order, note="table cell",
                                         anchor=valign))
            cy += row_h[ri]
        for _, box in edges.values():
            self.order += 1
            box.order = self.order
            boxes.append(box)
        return boxes, cy - y

    def _cell_borders(self, style: dict, x, y, w, h, table_box, edges: dict):
        """Adds a table cell's borders to edges, each (width, box) keyed by
        the edge it is on. Publisher collapses the borders of two
        cells that share an edge into one, the wider, centred on the edge,
        and draws those on the table's outside edges inside the table."""
        tx, ty, tw, th = table_box
        pen = to_px(style.get("x-pen-inset"), self.dpi) or 0.0
        if pen and w > 4 * pen and h > 4 * pen:
            # a shape's outline, which Publisher draws inside the shape
            x, y, w, h = x + pen, y + pen, w - 2 * pen, h - 2 * pen
            tx, ty, tw, th = tx + pen, ty + pen, tw - 2 * pen, th - 2 * pen
        for side in ("top", "bottom", "left", "right"):
            val = style.get(f"border-{side}") or style.get("border")
            if not val:
                continue
            toks = str(val).lower().split()
            if any(t in ("none", "hidden") for t in toks):
                continue
            color, width = None, None
            for t in toks:
                c = parse_color(t)
                if c is not None and color is None:
                    color = c
                bw = to_px(t, self.dpi)
                if bw is not None and width is None:
                    width = bw
            width = 1.0 if width is None else width
            if width <= 0:
                continue
            color = color if color is not None else RGBColor(0, 0, 0)
            if side in ("top", "bottom"):
                ey = y if side == "top" else y + h
                if abs(ey - ty) < 0.5:
                    top = ty
                elif abs(ey - (ty + th)) < 0.5:
                    top = ty + th - width
                else:
                    top = ey - width / 2
                left = max(x - width / 2, tx)
                rect = (left, top, min(x + w + width / 2, tx + tw) - left, width)
                key = ("h", round(ey * 2), round(x * 2), round((x + w) * 2))
            else:
                ex = x if side == "left" else x + w
                if abs(ex - tx) < 0.5:
                    left = tx
                elif abs(ex - (tx + tw)) < 0.5:
                    left = tx + tw - width
                else:
                    left = ex - width / 2
                top = max(y - width / 2, ty)
                rect = (left, top, width, min(y + h + width / 2, ty + th) - top)
                key = ("v", round(ex * 2), round(y * 2), round((y + h) * 2))
            seen = edges.get(key)
            if seen is not None and seen[0] >= width:
                continue
            box = Box(kind="rect", x=rect[0], y=rect[1], w=rect[2], h=rect[3],
                      fill=color, order=0, note="table cell border")
            edges[key] = (width, box)

    def _wrapper_padding(self, cell: Tag, sheet: StyleSheet, style: dict, cw: float) -> list:
        """Padding (left, top, right, bottom px) of the divs that wrap all of a
        cell's content, one inside the other."""
        total = [0.0, 0.0, 0.0, 0.0]
        el = cell
        while True:
            kids = [c for c in el.children
                    if isinstance(c, Tag) or (isinstance(c, NavigableString)
                                              and not isinstance(c, PreformattedString)
                                              and c.strip())]
            if len(kids) != 1 or not isinstance(kids[0], Tag) or kids[0].name != "div":
                return total
            el = kids[0]
            style = computed_style(el, sheet, style)
            for i, side in enumerate(("left", "top", "right", "bottom")):
                total[i] += to_px(style.get("padding-" + side), self.dpi, cw) or 0.0

    # -- segmentation: split a node's own content into document-order pieces --

    def _segments(self, el: Tag, sheet: StyleSheet, style: dict, stop_nodes: set) -> list:
        """[('text', [nodes], style) | ('img', node, style) | ('table', node, style)]"""
        segs: list = []
        buf: list = []
        buf_style = style

        def flush_buf():
            nonlocal buf
            if buf:
                segs.append(("text", buf, buf_style))
                buf = []

        def scan(node, cur_style):
            nonlocal buf_style
            for child in node.children:
                if isinstance(child, PreformattedString):
                    continue
                if isinstance(child, NavigableString):
                    if str(child).strip():
                        if not buf:
                            buf_style = cur_style
                        buf.append(child)
                    continue
                if not isinstance(child, Tag):
                    continue
                name = child.name.lower()
                if name in SKIP_TAGS or id(child) in stop_nodes:
                    continue
                cs = computed_style(child, sheet, cur_style)
                if name == "img":
                    flush_buf()
                    segs.append(("img", child, cs))
                    continue
                if name == "table":
                    flush_buf()
                    segs.append(("table", child, cs))
                    continue
                if (name in BLOCK_TAGS and child.find("img") is not None
                        and child.find(["table", *BLOCK_TAGS]) is None):
                    # a paragraph mixing pictures and text: lay it out as a line
                    flush_buf()
                    segs.append(("line", child, cs))
                    continue
                if child.find(["img", "table"]) is not None:
                    scan(child, cs)
                    continue
                if not buf:
                    buf_style = cur_style
                buf.append(child)
        scan(el, style)
        flush_buf()
        return segs

    # -- inline lines: pictures and text side by side -------------------------

    def _inline_items(self, node: Tag, sheet: StyleSheet, style: dict, out: list):
        for child in node.children:
            if isinstance(child, PreformattedString):
                continue
            if isinstance(child, NavigableString):
                text = _HTML_WS_RE.sub(" ", str(child))
                if text:
                    out.append(("text", text, style))
                continue
            if not isinstance(child, Tag):
                continue
            name = child.name.lower()
            if name in SKIP_TAGS:
                continue
            cs = computed_style(child, sheet, style)
            if name in ("img", "br"):
                out.append((name, child, cs))
                continue
            self._inline_items(child, sheet, cs, out)

    def _inline_text_box(self, img: Tag, sheet: StyleSheet, x, y, w, h,
                         boxes: list, fixed: bool) -> bool:
        """Places the text box an inline picture is of (see
        restore_vml_inline_text_boxes) as editable text, where stage 1 saw
        Publisher put it, else at the picture's own place (px). One wider
        than its frame runs past the frame's edge. Returns False if the
        picture isn't of a text box."""
        div = self.inline_boxes.get(img.get("x-textbox") or "")
        if div is None:
            return False
        key = pub_key(div.get_text())
        k = self.dpi / 72.0
        best = None
        for rec in self.pub_inline or []:
            if id(rec) in self.pub_claimed or rec["page"] not in (0, self.page):
                continue
            # the export can mangle characters the box's text ends with
            if rec["key"] != key and (len(key) < 12 or rec["key"][:20] != key[:20]):
                continue
            d = abs(rec["left"] * k - x) + abs(rec["top"] * k - y)
            if best is None or d < best[0]:
                best = (d, rec)
        rec = None
        if best:
            rec = best[1]
            self.pub_claimed.add(id(rec))
            x, y, w, h = rec["left"] * k, rec["top"] * k, rec["width"] * k, rec["height"] * k
            fixed = True
        elif self.pub_inline is not None:
            # stage 1 lists every inline text box Publisher shows
            self.warn(f"dropped an inline text box Publisher hides: {div.get_text(' ', strip=True)[:40]!r}")
            return True
        first = len(boxes)
        style = computed_style(div, sheet, {})
        outer, self.pub_box = self.pub_box, rec["box"] if rec else None
        self._emit_node(div, sheet, style, (x, y), w, h, boxes, set())
        self.pub_box = outer
        for b in boxes[first:]:
            b.fixed = b.fixed or fixed
            if b.kind != "text" or not b.paras:
                continue
            # without Publisher's positions, its space before each paragraph:
            # above the first, it moves the text down; it adds to the space
            # after the paragraph before the others
            for i, p in enumerate(b.paras):
                prec = p.pub if not p.pub_placed else None
                if not prec:
                    continue
                # Publisher's line spacing, which PowerPoint's single
                # spacing doesn't match
                multiple = pub_line_multiple(prec)
                sized = [r for r in p.runs if r.text.strip() and r.size_pt]
                if multiple and sized:
                    big = max(sized, key=lambda r: r.size_pt)
                    ratio = pub_single_spacing(big.font, big.bold, big.italic)
                    if ratio:
                        p.exact_line_pt = ratio * big.size_pt * multiple
                before = float(prec.get("spaceBefore") or 0.0)
                if before <= 0:
                    continue
                if i == 0:
                    b.y += before * k
                    b.h = max(b.h - before * k, 8.0)
                else:
                    b.paras[i - 1].space_after_pt += before
        if rec and rec["overflowing"]:
            # Publisher hides the paragraphs starting below the box. Stage 1
            # can't say where they are, so count their lines down from the top.
            bottom = y + h - (to_px(style.get("padding-bottom"), self.dpi) or 0.0)
            for b in boxes[first:]:
                if b.kind != "text":
                    continue
                at = b.y
                for i, p in enumerate(b.paras):
                    if at >= bottom - 1.0 and any(q.text().strip() for q in b.paras[i:]):
                        self.warn(f"dropped {len(b.paras) - i} paragraph(s) in a text box's "
                                  f"overflow, which Publisher hides: {p.text()[:40]!r}")
                        del b.paras[i:]
                        break
                    if p.keep_lines:
                        size_px = max((r.size_pt or 12.0) for r in p.runs) * k
                        line = p.exact_line_pt * k if p.exact_line_pt else \
                            size_px * (p.line_spacing or 1.2)
                        at += (p.text().count("\v") + 1) * line + p.space_after_pt * k
                    else:
                        at += self.estimate_height([p], b.w)
        return True

    def _line_boxes(self, para: Tag, sheet: StyleSheet, style: dict,
                    x, y, w, h, boxes: list, region=None) -> float:
        """Places a paragraph's pictures and text left to right, wrapping at the
        frame edge, as Publisher lays out a line with inline pictures. Spaces
        keep their width, and the export's tab runs snap to tab stops.
        Returns the height used."""
        def width(text, st, size_px, fam=None):
            """text's width in px, in its run's font, weight and style"""
            if fam is None:
                fam = (st.get("font-family") or "").split(",")[0].strip().strip("'\"")
            weight = str(st.get("font-weight") or "").lower()
            bold = weight in ("bold", "bolder") or (weight.isdigit() and int(weight) >= 600)
            italic = str(st.get("font-style") or "").lower() in ("italic", "oblique")
            return text_width_px(text, fam, size_px, bold, italic)
        items: list = []
        self._inline_items(para, sheet, style, items)
        gap = INLINE_PICTURE_GAP_PT * self.dpi / 72.0

        measured = []          # [kind, payload, style, width, height]
        after_space = True
        for kind, payload, st in items:
            if kind == "img":
                _, _, iw, ih = self._box_geometry(payload, st, (0.0, 0.0), w, h)
                iw = iw or 50.0
                ih = ih or iw
                measured.append(["img", payload, st, iw + 2 * gap, ih + 2 * gap])
                after_space = False
            elif kind == "br":
                measured.append(["br", None, st, 0.0, 0.0])
                after_space = True
            else:
                text = payload.lstrip(" ") if after_space else payload
                if not text:
                    continue
                after_space = text.endswith(" ")
                size_px = to_px(st.get("font-size", "12pt"), self.dpi) or 16.0
                for i, piece in enumerate(_WS_RUN_RE.split(text)):
                    if piece:
                        measured.append(["space" if i % 2 else "text", piece, st,
                                         width(piece, st, size_px), size_px * 1.2])

        # Publisher's own text for this line: real tabs and tab stops
        key = pub_key("".join(
            OBJECT_CHAR if m[0] == "img" else (m[1] or "") for m in measured))
        rec, placed = self.pub_find(key, region)
        if not placed and key.endswith(OBJECT_CHAR) and key.strip(OBJECT_CHAR):
            # Publisher leaves pictures in the overflow out of its text
            rec2, placed2 = self.pub_find(key.rstrip(OBJECT_CHAR), region)
            if placed2:
                rec, placed = rec2, placed2
                while measured and measured[-1][0] in ("img", "space", "br"):
                    measured.pop()
        pub_gaps, stops_px, pub_breaks = None, [], None
        if rec:
            h_words = []
            for m in measured:
                if m[0] == "img":
                    h_words.append(OBJECT_CHAR)
                elif m[0] == "text":
                    h_words.extend(m[1].split())
                elif m[0] == "space":
                    m.append(len(h_words))         # index of the word it precedes
            p_words, p_gaps = pub_words_and_gaps(rec["text"])
            if p_words == h_words:
                pub_gaps = p_gaps
                # the words Publisher starts its lines with, so text measured
                # without Publisher's own metrics breaks where Publisher's does
                starts = rec.get("lineStarts") or []
                if len(starts) > 1:
                    ends, at = [], 0           # where each word ends
                    for tok in _PUB_TOKEN_RE.split(rec["text"]):
                        at += len(tok)
                        if tok and not tok.isspace():
                            ends.append(at)
                    pub_breaks = {next(i for i, e in enumerate(ends) if e > s)
                                  for s in starts[1:] if 0 < s < at and ends and ends[-1] > s}
                    pub_breaks.discard(0)
            else:
                placed = False
                stops_px = sorted(t.get("pos", 0.0) * self.dpi / 72.0
                                  for t in rec.get("tabs") or [])

        align = str(style.get("text-align", "left")).lower()
        left_aligned = align not in ("center", "middle", "right")
        stop = TAB_STOP_PT * self.dpi / 72.0
        indent = to_px(style.get("text-indent"), self.dpi) or 0.0
        lines, cur, pos, wi = [], [], indent, 0
        for m in measured:
            if m[0] == "br":
                lines.append(cur)
                cur, pos = [], 0.0
                continue
            if m[0] == "space":
                size_px = m[4] / 1.2
                if pub_gaps is not None and m[5] in pub_gaps:
                    # Publisher's real whitespace: tabs go to their stops
                    space_w = width(" ", m[2], size_px)
                    end = pos
                    for ch in pub_gaps[m[5]]:
                        if ch == "\t":
                            end = next_tab_px(end, stops_px, stop)
                        elif ch not in "\v\n\r":
                            end += space_w
                    m[3] = end - pos
                elif left_aligned and m[1].count("\xa0") >= 2:
                    tabs = snap_to_tab(pos, m[3], width("\xa0", m[2], size_px), stop)
                    if tabs:
                        m[3] = (int(pos // stop) + tabs) * stop - pos
            if pub_breaks is not None:
                # break the lines where Publisher did
                if m[0] == "img":
                    if wi in pub_breaks and cur:
                        lines.append(cur)
                        cur, pos = [], 0.0
                    wi += 1
                elif m[0] == "text":
                    size_px = m[4] / 1.2
                    s = m[1]
                    while True:
                        cut = next((wd.start() for k, wd in enumerate(re.finditer(r"\S+", s))
                                    if (k or cur) and wi + k in pub_breaks), None)
                        if cut is None:
                            break
                        head = s[:cut]
                        wi += len(head.split())
                        if head.strip():
                            cur.append(["text", head, m[2], width(head, m[2], size_px), m[4]])
                        while cur and cur[-1][0] == "space":
                            cur.pop()
                        lines.append(cur)
                        cur, pos = [], 0.0
                        s = s[cut:]
                    if s != m[1]:
                        m = ["text", s, m[2], width(s, m[2], size_px), m[4]]
                    wi += len(s.split())
                elif m[0] == "space" and not cur:
                    continue
                cur.append(m)
                pos += m[3]
                continue
            while m[0] == "text" and pos + m[3] > w + 0.5 and " " in m[1].strip(" "):
                # too long for the line: break it between words, as Publisher does
                words = m[1].split(" ")
                size_px = m[4] / 1.2
                n = len(words) - 1
                while n > 0 and pos + width(" ".join(words[:n]), m[2], size_px) > w + 0.5:
                    n -= 1
                if n == 0 or not " ".join(words[:n]).strip():
                    if not cur:
                        break
                    lines.append(cur)
                    cur, pos = [], 0.0
                    m = [m[0], m[1].lstrip(" "), m[2],
                         width(m[1].lstrip(" "), m[2], size_px), m[4]]
                    continue
                head = " ".join(words[:n]) + " "
                tail = " ".join(words[n:]).lstrip(" ")
                cur.append(["text", head, m[2], width(head, m[2], size_px), m[4]])
                lines.append(cur)
                cur, pos = [], 0.0
                m = ["text", tail, m[2], width(tail, m[2], size_px), m[4]]
            if cur and pos + m[3] > w + 0.5:
                lines.append(cur)
                cur, pos = [], 0.0
                if m[0] == "space":
                    continue
            cur.append(m)
            pos += m[3]
        lines.append(cur)

        # Without stage 1's position, share the line's extra spacing above and
        # below it. With it, Publisher's line top is known and the extra
        # spacing falls below; the tallest picture sits its gap below the top.
        spacing = (pub_line_multiple(rec) if rec and pub_gaps is not None else None) \
            or self._line_multiple(style)
        # where Publisher puts the line, when stage 1 recorded it
        cy = rec["top"] * self.dpi / 72.0 if placed else y
        for li, line in enumerate(lines):
            if not line:
                continue
            lead = indent if li == 0 else 0.0
            total = sum(m[3] for m in line)
            if align in ("center", "middle"):
                cx = x + lead + max(w - lead - total, 0.0) / 2
            elif align == "right":
                cx = x + max(w - total, lead)
            else:
                cx = x + lead
            content_h = max(m[4] for m in line)
            extra = content_h * (spacing - 1.0)
            top = cy if placed else cy + extra / 2
            for kind, payload, st, iw, ih, *_ in line:
                if kind == "img" and self._inline_text_box(
                        payload, sheet, cx + gap, top + content_h - ih + gap,
                        iw - 2 * gap, ih - 2 * gap, boxes, placed):
                    pass
                elif kind == "img":
                    src = unquote(payload.get("src") or "")
                    if src:
                        self.order += 1
                        boxes.append(Box(kind="image", x=cx + gap,
                                         y=top + content_h - ih + gap,
                                         w=iw - 2 * gap, h=ih - 2 * gap,
                                         src=src, order=self.order,
                                         note=payload.get("alt") or "", fixed=placed))
                elif kind == "text" and payload.strip():
                    run = self._make_run(payload.strip(), st)
                    size_px = ih / 1.2
                    fam = run.font if run else None
                    skip = width(payload[:len(payload) - len(payload.lstrip())],
                                 st, size_px, fam)
                    vis = width(payload.strip(), st, size_px, fam)
                    self.order += 1
                    # slack so PowerPoint's own metrics don't wrap the text
                    boxes.append(Box(kind="text", x=cx + skip, y=top + content_h - ih,
                                     w=vis * 1.15 + 4.0, h=ih,
                                     paras=[Para(runs=[run])], order=self.order,
                                     note="inline text", fixed=placed))
                cx += iw
            cy += content_h + extra

        mb = to_px(style.get("margin-bottom"), self.dpi) or 0.0
        self.trailing = mb
        return cy - y + mb

    def _line_multiple(self, style: dict) -> float:
        """A paragraph's line-height as a multiple, such as 1.14 for 114%."""
        value = str(style.get("line-height") or "").strip()
        if value.endswith("%"):
            try:
                ratio = float(value[:-1]) / 100.0
            except ValueError:
                ratio = 1.0
        else:
            lh = to_px(value, self.dpi)
            fs = to_px(style.get("font-size", "12pt"), self.dpi)
            ratio = lh / fs if lh and fs else 1.0
        return min(max(ratio, 1.0), 3.0)

    # -- tree walk ----------------------------------------------------------

    def _shape_children(self, el: Tag, sheet: StyleSheet) -> list:
        out = []
        for child in el.find_all(True):
            if child.name in SKIP_TAGS:
                continue
            st = computed_style(child, sheet, {})
            if not (is_positioned(st) and (has_coords(st) or st.get("width"))):
                continue
            parent, nested = child.parent, False
            while parent is not None and parent is not el:
                if isinstance(parent, Tag):
                    ps = computed_style(parent, sheet, {})
                    if is_positioned(ps) and (has_coords(ps) or ps.get("width")):
                        nested = True
                        break
                parent = parent.parent
            if not nested:
                out.append(child)
        return out

    def walk(self, el: Tag, sheet: StyleSheet, inherited: dict,
             origin, cb_w, cb_h, boxes: list, depth=0):
        shape_children = self._shape_children(el, sheet)
        stop_nodes = {id(c) for c in shape_children}
        own_style = inherited if el.name == "body" else computed_style(el, sheet, inherited)

        start = len(boxes)
        self._emit_node(el, sheet, own_style, origin, cb_w, cb_h, boxes, stop_nodes)
        own = boxes[start:]
        under = self._shapes_before_picture(el, shape_children) if own else set()
        below = []

        for child in shape_children:
            cstyle = computed_style(child, sheet, own_style)
            x, y, w, h = self._box_geometry(child, cstyle, origin, cb_w, cb_h)
            new_w = w if w is not None else max(cb_w - (x - origin[0]), 1.0)
            new_h = h if h is not None else max(cb_h - (y - origin[1]), 1.0)
            first = len(boxes)
            if child.name.lower() == "img":
                src = unquote(child.get("src") or "")
                if src:
                    self.order += 1
                    adj = tuple(float(v) for v in str(cstyle.get("x-adj") or "").split())
                    boxes.append(Box(kind="image", x=x, y=y, w=new_w, h=new_h,
                                     src=src, order=self.order,
                                     note=child.get("alt") or "",
                                     geom=cstyle.get("x-geom") or "rect", adj=adj))
            else:
                self.walk(child, sheet, cstyle, (x, y), new_w, new_h, boxes, depth + 1)
            if id(child) in under:
                below += boxes[first:]
            # Publisher's stacking order is the z-index, not the order in the file
            try:
                z = int(str(cstyle.get("z-index") or "").strip())
            except ValueError:
                continue
            for b in boxes[first:]:
                if b.z is None:
                    b.z = z

        if below:
            # Publisher writes a group's picture after the members it exported
            # on their own, which lie beneath it: a group's first child is its
            # bottom one. Give those members the lower places in the order.
            orders = sorted(b.order for b in below + own)
            for b, o in zip(below + own, orders):
                b.order = o

    def _shapes_before_picture(self, el: Tag, shape_children: list) -> set:
        """ids of the shape children that are group members stacked beneath
        most of the shapes in el's own in-flow picture of the group, going
        by the order of the group's children in the VML. (The picture is one
        layer, so a member between its shapes goes where fewer are wrong: a
        white shape that hides part of one is usually above it.)"""
        kids = {id(c) for c in shape_children}
        order = getattr(self, "vml_order", None) or {}

        def places(node):
            imgs = [node] if node.name == "img" else node.find_all("img")
            return [order[s] for i in imgs for s in str(i.get("v:shapes") or "").split()
                    if s in order]

        pic = next((n for n in el.find_all("img")
                    if not any(id(p) in kids for p in n.parents)), None)
        # the picture's first shape is the group itself, ahead of its children
        drawn = places(pic)[1:] if pic is not None else []
        if not drawn:
            return set()
        under = set()
        for child in shape_children:
            mine = places(child)
            if mine and 2 * sum(1 for d in drawn if d > max(mine)) > len(drawn):
                under.add(id(child))
        return under

    def _emit_node(self, el: Tag, sheet: StyleSheet, style: dict,
                   origin, cb_w, cb_h, boxes: list, stop_nodes: set) -> float:
        """Adds the node's boxes and returns the height its content used."""
        x, y = origin
        w = cb_w if cb_w else 200.0
        h = cb_h if cb_h else 40.0

        # background / border of the container itself
        fill = background_color(style)
        border, border_w = None, 0.0
        for prop in ("border", "border-top", "border-left", "border-color", "border-width"):
            val = style.get(prop)
            if not val:
                continue
            for tok in str(val).split():
                c = parse_color(tok)
                if c is not None and border is None:
                    border = c
                bw = to_px(tok, self.dpi)
                if bw and bw > 0 and border_w == 0.0:
                    border_w = bw
        if border is not None and border_w == 0.0:
            border_w = 1.0
        # a rebuilt VML text box or picture says which outline shape it had
        adj = tuple(float(v) for v in str(style.get("x-adj") or "").split())
        if (fill is not None or border is not None) and w and h:
            # PowerPoint centres the line on the edge; Publisher drew it inside
            pen = min(to_px(style.get("x-pen-inset"), self.dpi) or 0.0, w / 4, h / 4)
            self.order += 1
            boxes.append(Box(kind="rect", x=x + pen, y=y + pen, w=w - 2 * pen, h=h - 2 * pen,
                             fill=fill,
                             line=border, line_w_px=border_w, order=self.order,
                             geom=style.get("x-geom") or "rect", adj=adj,
                             dash=style.get("x-dash"),
                             round_cap=style.get("x-cap") == "round"))

        if el.name.lower() == "img":
            src = unquote(el.get("src") or "")
            if src:
                self.order += 1
                boxes.append(Box(kind="image", x=x, y=y, w=w, h=h, src=src,
                                 order=self.order, note=el.get("alt") or "",
                                 geom=style.get("x-geom") or "rect", adj=adj))
            return h

        if el.name.lower() == "table":
            # a positioned table: keep its cell grid instead of flowing the cells
            tb, used = self._table_boxes(el, sheet, style, x, y, w, h, stop_nodes)
            boxes.extend(tb)
            return used

        pad_l = to_px(style.get("padding-left"), self.dpi, w) or 0.0
        pad_t = to_px(style.get("padding-top"), self.dpi, h) or 0.0
        pad_r = to_px(style.get("padding-right"), self.dpi, w) or 0.0
        content_x = x + pad_l
        content_w = max(w - pad_l - pad_r, 10.0)
        cursor = y + pad_t
        # text in columns: each line's place is in the first one's, and a
        # line that would reach into the bottom margin goes on in the next
        cols = max(int(float(style.get("x-columns") or 1)), 1)
        col_b = y + h - (to_px(style.get("padding-bottom"), self.dpi, h) or 0.0)
        gap = (to_px(style.get("x-column-gap"), self.dpi) or 0.0) if cols > 1 else 0.0
        col_w = max((content_w - (cols - 1) * gap) / cols, 10.0)

        segs = self._segments(el, sheet, style, stop_nodes)
        blank = self.pub_blank_frame(x, y, w, h)
        if blank:
            hidden = " ".join(el.get_text(" ").split())
            if hidden:
                self.warn(f"dropped a text box's text, all in its overflow, which "
                          f"Publisher hides: {hidden[:40]!r}")
            # its text goes, often in a layout table; pictures and borders stay
            self.hide_text += 1
        trailing = 0.0
        text_segs = [s for s in segs if s[0] == "text"]
        only_text = len(segs) == len(text_segs)

        overflowed = False
        for kind, payload, seg_style in segs:
            if overflowed and not has_coords(seg_style):
                # after the last text the frame shows: Publisher hides it
                if kind == "img":
                    self.warn("dropped a picture in a text box's overflow, which Publisher hides")
                continue
            if kind == "text":
                paras = self.collect_paras(payload, sheet, seg_style, stop_nodes,
                                           region=(x, y, x + w, y + h),
                                           text_area=(content_x, content_x + col_w))
                if not paras:
                    continue
                self.pub_overflowed = False
                where = self._pub_place(paras, bottom=y + h if cb_h else None)
                overflowed = bool(where) and self.pub_overflowed
                if where:
                    top, used = where
                    self.order += 1
                    # (text in columns fills them down to the bottom margin)
                    boxes.append(Box(kind="text", x=content_x, y=top, w=content_w,
                                     h=(max(col_b - top, 8.0) if cols > 1 else
                                        max(used, y + h - top if only_text else used, 8.0)),
                                     paras=paras, order=self.order, fixed=True,
                                     columns=cols, column_gap=gap))
                    cursor = top + used
                    trailing = 0.0
                    continue
                est = self.estimate_height(paras, content_w)
                if only_text and len(text_segs) == 1:
                    box_h = max(h - pad_t, 8.0)      # honour the declared frame
                else:
                    box_h = max(est, 8.0)
                self.order += 1
                boxes.append(Box(kind="text", x=content_x, y=cursor,
                                 w=content_w, h=box_h, paras=paras, order=self.order,
                                 columns=cols, column_gap=gap))
                cursor += est
                trailing = paras[-1].space_after_pt * self.dpi / 72.0
                continue

            if kind == "img":
                ix, iy, iw, ih = self._box_geometry(payload, seg_style,
                                                    (content_x, cursor), content_w, h)
                if not has_coords(seg_style):
                    ix, iy = content_x, cursor
                iw = iw or content_w
                ih = ih or (iw * 0.75)
                src = unquote(payload.get("src") or "")
                if self._inline_text_box(payload, sheet, ix, iy, iw, ih, boxes, False):
                    pass
                elif src:
                    self.order += 1
                    boxes.append(Box(kind="image", x=ix, y=iy, w=iw, h=ih, src=src,
                                     order=self.order, note=payload.get("alt") or ""))
                if not has_coords(seg_style):
                    cursor += ih
                    trailing = 0.0
                continue

            if kind == "line":
                cursor += self._line_boxes(payload, sheet, seg_style, content_x, cursor,
                                           content_w, h, boxes, region=(x, y, x + w, y + h))
                trailing = self.trailing
                continue

            # table
            tx, ty, tw, th = self._box_geometry(payload, seg_style,
                                                (content_x, cursor), content_w, h)
            if not has_coords(seg_style):
                tx, ty = content_x, cursor
            tb, used = self._table_boxes(payload, sheet, seg_style, tx, ty,
                                         tw or content_w, th, stop_nodes)
            boxes.extend(tb)
            if not has_coords(seg_style):
                cursor += used
                trailing = 0.0
        self.trailing = trailing
        if blank:
            self.hide_text -= 1
        return cursor - y


# ----------------------------------------------------------------------------
# Page detection
# ----------------------------------------------------------------------------

def _vml_pt(value, default=0.0) -> float:
    """A VML length in pt: '34.84pt', '0' or '8.68mm'."""
    m = re.match(r"\s*(-?[\d.]+)\s*(pt|in|px|mm|cm)?", str(value or ""))
    if not m:
        return default
    v = float(m.group(1))
    return {"in": v * 72.0, "px": v * 0.75, "mm": v * 72.0 / 25.4,
            "cm": v * 72.0 / 2.54}.get(m.group(2), v)


def _vml_fraction(value, default: float) -> float:
    """A VML fraction: '0.25', '25%' or '16384f' (65536ths)."""
    v = str(value or "").strip()
    try:
        if v.endswith("f"):
            return float(v[:-1]) / 65536.0
        if v.endswith("%"):
            return float(v[:-1]) / 100.0
        return float(v) if v else default
    except ValueError:
        return default


# VML shape type (#_x0000_tNN) -> outline: rectangles, rounded ones, ellipses
# and their callouts. Other shapes keep Publisher's picture.
_VML_TYPES = {"202": "rect", "1": "rect", "2": "roundRect", "3": "ellipse",
              "176": "roundRect", "61": "wedgeRectCallout",
              "62": "wedgeRoundRectCallout", "63": "wedgeEllipseCallout"}


def _vml_outline(shape: Tag) -> Optional[tuple]:
    """(geom, adjustments) for a VML shape's outline, or None if PowerPoint
    has no matching preset."""
    if shape.name in ("v:rect", "v:oval"):
        return ("rect" if shape.name == "v:rect" else "ellipse"), ()
    if shape.name == "v:roundrect":
        # arcsize: the corner radius as a share of the shorter side
        return "roundRect", (_vml_fraction(shape.get("arcsize"), 0.2),)
    m = re.match(r"#_x0000_t(\d+)$", str(shape.get("type") or ""))
    geom = _VML_TYPES.get(m.group(1)) if m else None
    if geom is None:
        return None
    adj = [float(v) for v in re.findall(r"-?\d+(?:\.\d+)?", str(shape.get("adj") or ""))]
    if m.group(1) == "2":
        return geom, ((adj[0] if adj else 3600.0) / 21600.0,)
    if m.group(1) == "176":
        return geom, (1 / 6,)
    if geom.endswith("Callout"):
        tx, ty = (adj + [1350.0, 25920.0][len(adj):])[:2]
        if 0 <= tx <= 21600 and 0 <= ty <= 21600:
            # the pointer ends inside the shape, so none shows
            return {"wedgeRectCallout": "rect", "wedgeEllipseCallout": "ellipse"}.get(
                geom, "roundRect"), ((1 / 6,) if geom == "wedgeRoundRectCallout" else ())
        return geom, (tx / 21600.0 - 0.5, ty / 21600.0 - 0.5) + (
            (1 / 6,) if geom == "wedgeRoundRectCallout" else ())
    return geom, ()


def _vml_text_area(shape: Tag, shapetype: Optional[Tag], geom: str, adj: tuple,
                   width: float, height: float) -> tuple:
    """How far (left, top, right, bottom pt) a shape's text area sits inside
    its outline: the shape type's textboxrect, in its coordsize. Formula
    rectangles of rounded shapes are worked out from the corner radius, and
    the ellipse's from its preset."""
    path = (shape.find("v:path") or (shapetype.find("v:path") if shapetype else None))
    rect = str(path.get("textboxrect") or "") if path else ""
    nums = rect.split(";")[0].split(",")
    if len(nums) == 4 and all(re.fullmatch(r"-?\d+", n.strip()) for n in nums):
        size = str((shapetype or shape).get("coordsize") or "21600,21600").split(",")
        try:
            cw, ch = float(size[0]), float(size[-1])
        except ValueError:
            cw = ch = 21600.0
        l, t, r, b = (float(n) for n in nums)
        return (l / cw * width, t / ch * height, (cw - r) / cw * width, (ch - b) / ch * height)
    if geom == "roundRect" and adj:
        k = (1 - 0.5 ** 0.5) * adj[0] * min(width, height)   # where the corner arc is at 45 degrees
        return (k, k, k, k)
    if geom.startswith(("ellipse", "wedgeEllipse")):
        return (0.1464 * width, 0.1464 * height, 0.1464 * width, 0.1464 * height)
    return (0.0, 0.0, 0.0, 0.0)


DOWNLEVEL_RE = re.compile(rb"<!\[if\s+([^\]]*)\]>|<!\[endif\]>", re.I)


def drop_downlevel_fallbacks(raw: bytes) -> bytes:
    """Publisher can write a shape inline as <![if mso]>VML<![endif]> followed by
    <![if !mso]><img><![endif]>, the <img> being a picture of the same shape for
    other browsers. lxml drops the markers and keeps both, so the shape's text
    would arrive twice: once as text, once as a picture. Removes each !mso branch
    that directly follows an mso branch. Branches nest (<![if RotText]>)."""
    out, pos, depth, mso_end = [], 0, 0, None
    stack = []
    for m in DOWNLEVEL_RE.finditer(raw):
        if m.group(1) is not None:
            cond = m.group(1).strip().lower()
            if depth == 0 and cond == b"!mso" and mso_end is not None \
                    and not raw[mso_end:m.start()].strip():
                out.append(raw[pos:m.start()])
                pos = None
            stack.append((cond, m.start()))
            depth += 1
            continue
        if not stack:
            continue
        cond, _ = stack.pop()
        depth -= 1
        if depth:
            continue
        if pos is None:
            pos = m.end()               # resume after the dropped !mso branch
            mso_end = None
        else:
            mso_end = m.end() if cond == b"mso" else None
    if pos is None:
        return raw
    out.append(raw[pos:])
    return b"".join(out)


def restore_vml_pictures(soup: BeautifulSoup) -> int:
    """A picture cropped to a shape (a rounded rectangle, say) can come only in
    the VML, with no <img> for other browsers. Adds a positioned <img> for each,
    saying which shape crops it. Returns how many were added."""
    shown = {el.get("v:shapes") for el in soup.find_all(attrs={"v:shapes": True})}
    count = 0
    for c in soup.find_all(string=lambda t: isinstance(t, Comment) and "vml" in t[:30]):
        vml = BeautifulSoup(str(c), "html.parser")
        for shape in vml.find_all(["v:shape", "v:rect", "v:roundrect", "v:oval"]):
            data = shape.find("v:imagedata", recursive=False)
            decl = parse_decls(shape.get("style", ""))
            if (data is None or not data.get("src") or shape.get("id") in shown
                    or shape.find_parent("v:group") or decl.get("rotation")
                    or str(decl.get("position", "")).lower() != "absolute"):
                continue
            geom, adj = _vml_outline(shape) or ("rect", ())
            styles = ["position:absolute"] + [
                f"{k}:{_vml_pt(decl.get(k)):.2f}pt" for k in ("left", "top", "width", "height")]
            if decl.get("z-index"):
                styles.append(f"z-index:{decl['z-index']}")
            if geom != "rect":
                styles.append(f"x-geom:{geom}")
                if adj:
                    styles.append("x-adj:" + " ".join(f"{v:.5f}" for v in adj))
            img = soup.new_tag("img", src=data["src"], style=";".join(styles),
                               alt=shape.get("alt") or "")
            img["v:shapes"] = shape.get("id") or ""
            c.insert_after(img)
            shown.add(shape.get("id"))
            count += 1
    return count


def _vml_box_styles(shape: Tag, decl: dict, box: Tag, outline: tuple,
                    shapetype: Optional[Tag]) -> list:
    """A VML text box's size, padding, fill and outline, as CSS declarations."""
    geom, adj = outline
    inset = [_vml_pt(v, 2.88) for v in (box.get("inset") or "").split(",")]
    inset += [2.88] * (4 - len(inset))          # left, top, right, bottom
    inner = box.find("div")
    pad = parse_decls(inner.get("style", "")) if inner else {}
    width, height = _vml_pt(decl.get("width")), _vml_pt(decl.get("height"))
    # the shape's own text area inside its outline
    area = _vml_text_area(shape, shapetype, geom, adj, width, height)
    styles = [
        f"width:{width:.2f}pt",
        f"height:{height:.2f}pt",
        f"padding-left:{area[0] + inset[0] + _vml_pt(pad.get('padding-left')):.2f}pt",
        f"padding-top:{area[1] + inset[1] + _vml_pt(pad.get('padding-top')):.2f}pt",
        f"padding-right:{area[2] + inset[2] + _vml_pt(pad.get('padding-right')):.2f}pt",
        f"padding-bottom:{area[3] + inset[3] + _vml_pt(pad.get('padding-bottom')):.2f}pt",
    ]
    cols = parse_decls(box.get("style", ""))
    if int(_vml_pt(cols.get("mso-column-count"), 1.0)) > 1:
        styles += [f"x-columns:{int(_vml_pt(cols.get('mso-column-count')))}",
                   f"x-column-gap:{_vml_pt(cols.get('mso-column-margin')):.2f}pt"]
    # VML shapes are filled unless they say otherwise, white by default
    fill = str(shape.get("fillcolor") or "white").split() or ["white"]
    fill_el = shape.find("v:fill", recursive=False)
    if (str(shape.get("filled", "t")).lower() not in ("f", "false")
            and str((fill_el or {}).get("on", "t")).lower() not in ("f", "false")):
        styles.append(f"background-color:{fill[0]}")
    stroke = shape.find("v:stroke", recursive=False)
    if (str(shape.get("stroked", "t")).lower() not in ("f", "false")
            and str((stroke or {}).get("on", "t")).lower() not in ("f", "false")):
        color = (str(shape.get("strokecolor") or "black").split() or ["black"])[0]
        styles.append(f"border:{_vml_pt(shape.get('strokeweight'), 0.75):.2f}pt solid {color}")
        if _vml_inset_pen(shape):
            styles.append(f"x-pen-inset:{_vml_inset_pen(shape):.2f}pt")
        dash = str((stroke or {}).get("dashstyle") or "").lower()
        if dash and dash != "solid":
            styles.append(f"x-dash:{dash}")
        if str((stroke or {}).get("endcap") or "").lower() == "round":
            styles.append("x-cap:round")
    elif stroke is not None:
        # a text box's border set side by side: drawn when all four sides
        # have one, a weight of 0 being a hairline
        sides = [stroke.find(f"o:{side}", recursive=False)
                 for side in ("left", "top", "right", "bottom")]
        if all(sd is not None and str(sd.get("on", "")).lower() in ("t", "true")
               for sd in sides):
            color = (str(sides[0].get("color") or "black").split() or ["black"])[0]
            weight = max(_vml_pt(sides[0].get("weight"), 0.75), 0.5)
            styles.append(f"border:{weight:.2f}pt solid {color}")
    if geom != "rect":
        styles.append(f"x-geom:{geom}")
    if adj:
        styles.append("x-adj:" + " ".join(f"{v:.5f}" for v in adj))
    return styles


def _vml_flat_box(shape: Tag) -> bool:
    """True if a shape draws only a flat fill and plain outline: no picture,
    gradient, pattern or texture fill, no shadow, no second stroke colour or
    line style (BorderArt is exported as a white stroke with one), and any
    border set side by side alike on all four sides."""
    fill = shape.find("v:fill", recursive=False)
    stroke = shape.find("v:stroke", recursive=False)
    shadow = shape.find("v:shadow", recursive=False)
    sides = {(str(sd.get("on", "")).lower(), sd.get("weight"), sd.get("color"))
             for sd in (stroke.find(f"o:{side}", recursive=False) if stroke is not None else None
                        for side in ("left", "top", "right", "bottom"))
             if sd is not None}
    return (shape.find("v:imagedata") is None
            and (stroke is None or set(stroke.attrs) <= {"dashstyle", "joinstyle", "endcap"})
            and len(sides) <= 1
            and (fill is None or (not fill.get("src")
                                  and str(fill.get("type") or "solid").lower() == "solid"))
            and (shadow is None or str(shadow.get("on", "f")).lower() not in ("t", "true")))


def restore_vml_text_boxes(soup: BeautifulSoup) -> int:
    """Publisher exports some text boxes (filled ones, for instance) as a
    picture of the text, keeping the real text only in the VML inside an
    <!--[if gte vml 1]> comment. Swaps each such picture for a positioned div
    holding that text, so it stays editable, with the shape's outline behind
    it; an empty one too, if its picture shows only a flat fill and outline.
    Shapes inside a VML group, rotated or WordArt shapes and outlines
    PowerPoint has no preset for keep their picture; inline ones are left to
    restore_vml_inline_text_boxes. Returns how many were swapped."""
    found, text_rects = {}, {}
    for c in soup.find_all(string=lambda t: isinstance(t, Comment) and "vml" in t[:30]):
        vml = BeautifulSoup(str(c), "html.parser")
        for shape in vml.find_all(["v:shape", "v:rect", "v:roundrect", "v:oval"]):
            box = shape.find("v:textbox", recursive=False)
            if (box is None or shape.find_parent("v:group") or shape.find("v:textpath")
                    or not (box.get_text(strip=True) or _vml_flat_box(shape))):
                continue
            decl = parse_decls(shape.get("style", ""))
            if (decl.get("rotation") or "layout-flow" in str(box.get("style", ""))
                    or str(decl.get("position", "")).lower() != "absolute"):
                continue
            outline = _vml_outline(shape)
            if outline is None:
                continue
            found[shape.get("id")] = (shape, decl, box, outline)
        for st in vml.find_all("v:shapetype"):
            text_rects[st.get("id")] = st
    count = 0
    for img in soup.find_all("img"):
        hit = found.get(img.get("v:shapes"))
        if not hit:
            continue
        shape, decl, box, outline = hit
        styles = [
            "position:absolute",
            f"left:{_vml_pt(decl.get('left')):.2f}pt",
            f"top:{_vml_pt(decl.get('top')):.2f}pt",
        ] + _vml_box_styles(shape, decl, box, outline,
                            text_rects.get(str(shape.get("type") or "").lstrip("#")))
        if decl.get("z-index"):
            styles.append(f"z-index:{decl['z-index']}")
        div = soup.new_tag("div", style=";".join(styles))
        inner = box.find("div")
        for child in list((inner or box).children):
            div.append(child.extract())
        # the picture sits in a positioned span of its own
        target = img.parent if (img.parent and img.parent.name == "span"
                                and len(img.parent.find_all(True)) == 1) else img
        target.replace_with(div)
        count += 1
    return count


VML_SHAPES = ["v:shape", "v:rect", "v:roundrect", "v:oval"]


def _vml_inset_pen(shape: Tag) -> float:
    """Half the outline weight (pt) of a shape whose outline Publisher draws
    inside its edge (insetpen), else 0."""
    stroke = shape.find("v:stroke", recursive=False)
    if (str(shape.get("insetpen", "")).lower() not in ("t", "true")
            or str(shape.get("stroked", "t")).lower() in ("f", "false")
            or str((stroke or {}).get("on", "t")).lower() in ("f", "false")):
        return 0.0
    return _vml_pt(shape.get("strokeweight"), 0.75) / 2


def mark_inset_pens(soup: BeautifulSoup) -> int:
    """Publisher draws a shape's outline inside its edge, but its HTML gives
    the table cell standing in for the shape a border centred on that edge,
    half the outline wider all round. Marks each such cell with x-pen-inset,
    how far in to draw it. Returns how many were marked."""
    inset = {}
    for c in soup.find_all(string=lambda t: isinstance(t, Comment) and "vml" in t[:30]):
        for shape in BeautifulSoup(str(c), "html.parser").find_all(VML_SHAPES):
            if not shape.find_parent("v:group") and _vml_inset_pen(shape):
                inset[shape.get("id")] = _vml_inset_pen(shape)
    count = 0
    for div in soup.find_all(attrs={"v:shape": True}):
        half = inset.get(div.get("v:shape"))
        td = div.find_parent("td")
        if half and td is not None and "border" in str(td.get("style", "")):
            td["style"] = str(td.get("style", "")).rstrip("; ") + f";x-pen-inset:{half:.2f}pt"
            count += 1
    return count


def _vml_group_parts(group: Tag, pics: list) -> Optional[list]:
    """A VML group's pictures and text boxes, each with its place in pt:
    ("img", shape, rect, pic) with stage 1's picture for it, or ("text",
    shape, rect, textbox, outline). None if the group holds anything else,
    or a picture stage 1 didn't save."""
    decl = parse_decls(group.get("style", ""))
    if decl.get("rotation") or str(decl.get("position", "")).lower() != "absolute":
        return None
    gl, gt = _vml_pt(decl.get("left")), _vml_pt(decl.get("top"))
    gw, gh = _vml_pt(decl.get("width")), _vml_pt(decl.get("height"))
    try:
        ox, oy = (float(v) for v in str(group.get("coordorigin") or "0,0").split(","))
        cw, ch = (float(v) for v in str(group.get("coordsize") or "1000,1000").split(","))
    except ValueError:
        return None
    if not (gw and gh and cw and ch):
        return None
    parts = []
    for child in group.find_all(True, recursive=False):
        if child.name not in VML_SHAPES:
            if child.name in ("o:lock", "v:fill", "v:stroke", "v:shadow", "v:shapetype"):
                continue
            return None
        cd = parse_decls(child.get("style", ""))
        if cd.get("rotation") or child.find("v:textpath"):
            return None
        try:
            rect = (gl + (float(cd["left"]) - ox) * gw / cw,
                    gt + (float(cd["top"]) - oy) * gh / ch,
                    float(cd["width"]) * gw / cw, float(cd["height"]) * gh / ch)
        except (KeyError, ValueError):
            return None
        box = child.find("v:textbox", recursive=False)
        data = child.find("v:imagedata", recursive=False)
        if box is not None and box.get_text(strip=True):
            outline = _vml_outline(child)
            if outline is None or "layout-flow" in str(box.get("style", "")):
                return None
            parts.append(("text", child, rect, box, outline))
        elif data is not None and data.get("src"):
            fits = [(misfit, pic) for pic in pics
                    for misfit in [_place_misfit(*rect, pic)] if misfit is not None]
            if not fits:
                return None
            parts.append(("img", child, rect, min(fits, key=lambda f: f[0])[1]))
        elif _vml_draws(child):
            # a shape with nothing in it, drawn as its outline and fill
            outline = _vml_outline(child)
            if outline is None:
                return None
            parts.append(("text", child, rect, box, outline))
        # an empty frame draws nothing in the group's picture either
    return parts


def _vml_draws(shape: Tag) -> bool:
    """Whether a VML shape is filled or outlined (VML's default is both)."""
    def on(attr, el):
        return (str(shape.get(attr, "t")).lower() not in ("f", "false")
                and str((el or {}).get("on", "t")).lower() not in ("f", "false"))
    return (on("filled", shape.find("v:fill", recursive=False))
            or on("stroked", shape.find("v:stroke", recursive=False)))


def restore_vml_groups(soup: BeautifulSoup, pics: Optional[list], page: int) -> int:
    """Publisher exports a group, a picture with its caption say, as one
    picture of the lot. Where the group holds only pictures and text boxes,
    and stage 1 saved each picture (pics, from load_pub_pictures), swaps it
    for those pictures and the text boxes' editable text, each in its place.
    Other groups keep their picture. Returns how many were swapped."""
    if not pics:
        return 0
    pics = [p for p in pics if p["page"] in (0, page)]
    text_rects, found = {}, {}
    for c in soup.find_all(string=lambda t: isinstance(t, Comment) and "vml" in t[:30]):
        vml = BeautifulSoup(str(c), "html.parser")
        for st in vml.find_all("v:shapetype"):
            text_rects[st.get("id")] = st
        for group in vml.find_all("v:group"):
            if group.find_parent("v:group"):
                continue
            parts = _vml_group_parts(group, pics)
            if (parts and any(p[0] == "img" for p in parts)
                    and any(p[0] == "text" and p[3] is not None and p[3].get_text(strip=True)
                            for p in parts)):
                found[group.get("id")] = (parts, parse_decls(group.get("style", "")).get("z-index"),
                                          group)
    count = 0
    for img in soup.find_all("img"):
        ids = str(img.get("v:shapes") or "").split()
        hit = found.get(ids[0]) if ids else None
        if not hit:
            continue
        parts, z, group = hit
        target = img.parent if (img.parent and img.parent.name == "span"
                                and len(img.parent.find_all(True)) == 1) else img
        # the group's place is measured outside any span positioning its picture
        anchor = target
        while (anchor.parent is not None and anchor.parent.name == "span" and "absolute"
               in str(parse_decls(anchor.parent.get("style", "")).get("position", ""))):
            anchor = anchor.parent
        # members stacked above a picture the span keeps for another member go after it
        order = [c.get("id") for c in group.find_all(True, recursive=False)]
        kept = [order.index(i) for other in anchor.find_all("img") if other is not img
                for i in str(other.get("v:shapes") or "").split() if i in order]
        above_from = min(kept) if anchor is not target and kept else len(order)
        after = anchor
        for part in parts:
            if anchor is not target and part[1].get("id") in order \
                    and order.index(part[1].get("id")) in kept:
                continue
            left, top, width, height = part[2]
            styles = ["position:absolute", f"left:{left:.2f}pt", f"top:{top:.2f}pt"]
            if part[0] == "img":
                styles += [f"width:{width:.2f}pt", f"height:{height:.2f}pt"]
                if z:
                    styles.append(f"z-index:{z}")
                el = soup.new_tag("img", src=os.path.basename(part[3]["png"]),
                                  style=";".join(styles))
            else:
                _, shape, _, box, outline = part
                if box is None:
                    box = soup.new_tag("v:textbox")
                styles += _vml_box_styles(
                    shape, {"width": f"{width:.2f}pt", "height": f"{height:.2f}pt"}, box,
                    outline, text_rects.get(str(shape.get("type") or "").lstrip("#")))
                if z:
                    styles.append(f"z-index:{z}")
                el = soup.new_tag("div", style=";".join(styles))
                inner = box.find("div")
                for child in list((inner or box).children):
                    el.append(child.extract())
            if part[1].get("id") in order and order.index(part[1].get("id")) > above_from:
                after.insert_after(el)
                after = el
            else:
                anchor.insert_before(el)
        target.extract()
        count += 1
    return count


def vml_shape_order(soup: BeautifulSoup) -> dict:
    """Each VML shape's id and its place in the file, which within a group
    is its stacking order, bottom first."""
    order = {}
    for c in soup.find_all(string=lambda t: isinstance(t, Comment) and "vml" in t[:30]):
        for m in re.finditer(r"<v:(?!shapetype)\w+\s[^>]*?\bid=\"([^\"]+)\"", str(c)):
            order.setdefault(m[1], len(order))
    return order


def restore_vml_inline_text_boxes(soup: BeautifulSoup) -> dict:
    """Publisher exports a text box placed in another one's text as a picture
    of it, right after the VML that keeps its real text. Marks each such
    picture with x-textbox, the key of a div holding that text with the box's
    size, padding and outline, for the converter to put in the picture's
    place. A text box inside such a box's text becomes an <img> of its size
    in the div, marked the same way. Rotated or WordArt shapes and outlines
    PowerPoint has no preset for keep their picture. Returns the divs by key."""
    divs: dict = {}
    shapetypes: dict = {}
    comments = [c for c in soup.find_all(
        string=lambda t: isinstance(t, Comment) and "vml" in t[:30])]
    parsed = [(c, BeautifulSoup(str(c), "html.parser")) for c in comments]
    for _, vml in parsed:
        for st in vml.find_all("v:shapetype"):
            shapetypes[st.get("id")] = st

    def rebuild(shape: Tag) -> Optional[Tag]:
        box = shape.find("v:textbox", recursive=False)
        if box is None or not box.get_text(strip=True) or shape.find("v:textpath"):
            return None
        decl = parse_decls(shape.get("style", ""))
        if (decl.get("rotation") or "layout-flow" in str(box.get("style", ""))
                or str(decl.get("position", "")).lower() == "absolute"):
            return None
        outline = _vml_outline(shape)
        if outline is None:
            return None
        styles = _vml_box_styles(shape, decl, box, outline,
                                 shapetypes.get(str(shape.get("type") or "").lstrip("#")))
        div = soup.new_tag("div", style=";".join(styles))
        content = box.find("div") or box
        # the text boxes in this one's text, outermost first
        for inner in [s for s in content.find_all(VML_SHAPES)
                      if s.find_parent(VML_SHAPES) is shape]:
            d = parse_decls(inner.get("style", ""))
            size = f"width:{_vml_pt(d.get('width')):.2f}pt;height:{_vml_pt(d.get('height')):.2f}pt"
            sub = rebuild(inner)
            data = inner.find("v:imagedata", recursive=False)
            # its picture for other browsers follows it, past <![if !vml]>
            nxt = inner.next_sibling
            while isinstance(nxt, PreformattedString) or (
                    isinstance(nxt, NavigableString) and not nxt.strip()):
                nxt = nxt.next_sibling
            if isinstance(nxt, Tag) and nxt.name == "img":
                nxt.decompose()
            if sub is not None:
                img = soup.new_tag("img", style=size)
                img["x-textbox"] = key = f"textbox{len(divs)}"
                divs[key] = sub
            elif data is not None and data.get("src"):
                img = soup.new_tag("img", src=data["src"], style=size,
                                   alt=inner.get("alt") or "")
            else:
                inner.decompose()
                continue
            inner.replace_with(img)
        for child in list(content.children):
            div.append(child.extract())
        return div

    for c, vml in parsed:
        tops = [s for s in vml.find_all(VML_SHAPES) if s.find_parent(VML_SHAPES) is None]
        if len(tops) != 1:
            continue
        # the picture for other browsers follows the VML
        nxt = c.next_sibling
        while isinstance(nxt, NavigableString) and not isinstance(nxt, Comment) \
                and not nxt.strip():
            nxt = nxt.next_sibling
        if not isinstance(nxt, Tag) or nxt.name != "img" or nxt.get("v:shapes"):
            continue
        div = rebuild(tops[0])
        if div is not None:
            nxt["x-textbox"] = key = f"textbox{len(divs)}"
            divs[key] = div
    return divs


def find_page_containers(soup: BeautifulSoup, sheet: StyleSheet, dpi: float,
                         default_h: Optional[float] = None,
                         default_w: Optional[float] = None):
    """Publisher wraps each page in a sized, positioned div. Find them.
    Publisher sometimes writes a malformed height ('10.-1737in') or width
    ('4.-2040in') on that div; default_h and default_w (px) stand in."""
    body = soup.body or soup
    candidates = []
    for el in body.find_all(["div", "table", "section"]):
        st = computed_style(el, sheet, {})
        w = to_px(st.get("width"), dpi)
        h = to_px(st.get("height"), dpi)
        # Publisher's own wrapper: a positioned div straight under <body>
        wrapper = el.name == "div" and el.parent is body and is_positioned(st)
        if not h and default_h and wrapper and st.get("height"):
            h = default_h
        # without its width the wrapper is skipped, and a large table on
        # the page is taken for the page instead
        if not w and default_w and wrapper and st.get("width"):
            w = default_w
        if not w or not h or ((w < 200 or h < 200) and not wrapper):
            continue
        # must not be nested inside another candidate
        candidates.append((el, w, h, st))

    top = []
    for el, w, h, st in candidates:
        if not any(other is not el and other in el.parents for other, _, _, _ in candidates):
            top.append((el, w, h, st))
    return top


def read_page_size(export_dir: str, base: str) -> Optional[tuple]:
    """The real page size in inches, recorded by stage 1 as <base>_pagesize.txt."""
    path = os.path.join(export_dir, base + "_pagesize.txt")
    try:
        with open(path, encoding="ascii") as fh:
            w, h = (float(v) for v in fh.read().split()[:2])
    except (OSError, ValueError):
        return None
    return (w, h) if w > 0 and h > 0 else None


def read_page_spreads(export_dir: str, base: str) -> Optional[list]:
    """Each page's (spread, x offset, y offset in pt) within its two-page
    spread, recorded by stage 1 as <base>_spreads.txt."""
    path = os.path.join(export_dir, base + "_spreads.txt")
    try:
        with open(path, encoding="ascii") as fh:
            rows = [line.split() for line in fh if line.strip()]
        return [(int(r[0]), float(r[1]), float(r[2])) for r in rows]
    except (OSError, ValueError, IndexError):
        return None


def share_spread_shapes(slide_pages: list, spreads: list, dpi: float) -> int:
    """Publisher keeps a shape that crosses the fold of a two-page spread on
    one page, and exports it only there. Put a copy, moved by the pages'
    offsets in the spread, on each other page of the spread it reaches, so
    that page shows its part. Returns the number of copies."""
    k = dpi / 72.0
    added = []
    for p, (spread, px, py) in enumerate(spreads):
        for q, (other, qx, qy) in enumerate(spreads):
            if q == p or other != spread:
                continue
            dx, dy = (px - qx) * k, (py - qy) * k
            qw, qh = slide_pages[q][0], slide_pages[q][1]
            for b in slide_pages[p][2]:
                x, y = b.x + dx, b.y + dy
                # ignore a hairline overlap: shapes drawn to the fold
                if x + (b.w or 0) > 1 and x < qw - 1 and y + (b.h or 0) > 1 and y < qh - 1:
                    added.append((q, dataclasses.replace(b, x=x, y=y)))
    for q, b in added:
        slide_pages[q][2].append(b)
    return len(added)


def bounding_page(boxes: list, dpi: float, fallback_in):
    if not boxes:
        return fallback_in[0] * dpi, fallback_in[1] * dpi
    max_x = max(b.x + (b.w or 0) for b in boxes)
    max_y = max(b.y + (b.h or 0) for b in boxes)
    # never smaller than the fallback page: an unpositioned publication is
    # almost always a normal page that Publisher exported as flowed HTML
    return max(max_x, fallback_in[0] * dpi), max(max_y, fallback_in[1] * dpi)


# ----------------------------------------------------------------------------
# Image resolution
# ----------------------------------------------------------------------------

class ImageResolver:
    def __init__(self, html_dir: str, export_dir: str, warn):
        self.html_dir = html_dir
        self.export_dir = export_dir
        self.warn = warn
        self.index = {}
        self.used = set()
        self.converted = {}
        for root, _dirs, files in os.walk(export_dir):
            for f in files:
                self.index.setdefault(f.lower(), os.path.join(root, f))

    def resolve(self, src: str) -> Optional[str]:
        if not src or src.lower().startswith(("http://", "https://", "data:")):
            return None
        src = src.replace("\\", "/").split("?")[0].split("#")[0]
        cand = os.path.normpath(os.path.join(self.html_dir, src))
        if os.path.isfile(cand):
            return self._usable(cand)
        base = os.path.basename(src).lower()
        hit = self.index.get(base)
        if hit:
            return self._usable(hit)
        self.warn(f"image not found on disk: {src}")
        return None

    def _usable(self, path: str) -> Optional[str]:
        ext = os.path.splitext(path)[1].lower()
        if ext in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"):
            if ext in (".bmp", ".tif", ".tiff"):
                return self._to_png(path)
            self.used.add(os.path.normcase(os.path.abspath(path)))
            return path
        if ext in (".wmf", ".emf", ".webp", ".ico"):
            return self._to_png(path)
        self.warn(f"unsupported image type skipped: {os.path.basename(path)}")
        return None

    def _to_png(self, path: str) -> Optional[str]:
        if path in self.converted:
            return self.converted[path]
        if not HAVE_PIL:
            self.warn(f"cannot convert {os.path.basename(path)} without Pillow")
            return None
        out = os.path.join(self.export_dir, "_pptx_converted")
        os.makedirs(out, exist_ok=True)
        dest = os.path.join(out, os.path.splitext(os.path.basename(path))[0] + ".png")
        try:
            with Image.open(path) as im:
                im.convert("RGBA" if im.mode in ("RGBA", "LA", "P") else "RGB").save(dest)
            self.converted[path] = dest
            self.used.add(os.path.normcase(os.path.abspath(path)))
            return dest
        except Exception as exc:                      # noqa: BLE001
            self.warn(f"could not convert {os.path.basename(path)}: {exc}")
            self.converted[path] = None
            return None


def aspect_of(path: str) -> Optional[float]:
    if not HAVE_PIL:
        return None
    try:
        with Image.open(path) as im:
            return im.width / im.height if im.height else None
    except Exception:                                  # noqa: BLE001
        return None


def hires_candidates(export_dir: str, base: str, page_no: int) -> list:
    pat = os.path.join(export_dir, f"{base}_page_{page_no:03d}_image_*.png")
    return sorted(globmod.glob(pat))


def load_pub_pictures(export_dir: str, base: str) -> Optional[list]:
    """Stage 1's record of where each exported PNG's picture sits, in pt:
    dicts with file, page (0 for a master page), left, top, width, height
    and rotation. None if stage 1 didn't record them."""
    path = os.path.join(export_dir, base + "_pictures.json")
    try:
        with open(path, encoding="utf-8-sig") as fh:
            data = json.load(fh)
    except (OSError, ValueError):
        return None
    pics = data.get("pictures") if isinstance(data, dict) else None
    if isinstance(pics, dict):
        # PowerShell writes a list of one as the item itself
        pics = [pics]
    out = []
    for pic in pics or []:
        try:
            png = os.path.join(export_dir, str(pic["file"]))
            if os.path.isfile(png):
                out.append({"png": png, "page": int(pic.get("page", 0)),
                             "left": float(pic["left"]), "top": float(pic["top"]),
                             "width": float(pic["width"]), "height": float(pic["height"]),
                             "rotation": float(pic.get("rotation") or 0.0)})
        except (KeyError, TypeError, ValueError):
            continue
    return out


# ----------------------------------------------------------------------------
# PPTX emission
# ----------------------------------------------------------------------------

def px_to_emu(px: float, dpi: float) -> int:
    return int(round(px / dpi * EMU_PER_INCH))


def _ppt_size(run, para) -> Optional[float]:
    """The size PowerPoint gets for a run: a superscript or subscript at the
    size it shrinks from, which is no more than the paragraph's other text
    (the export can leave such a run at the paragraph's default size)."""
    if not run.script or not run.size_pt:
        return run.size_pt
    others = [r.size_pt for r in para.runs if not r.script and r.size_pt and r.text.strip()]
    size = run.size_pt / SCRIPT_SCALE
    return min(size, max(others)) if others else size


def _line_sizes(para) -> set:
    """The largest font size on each line of a paragraph, spaces included;
    None for a run with no size."""
    return set(_line_size_list(para))


def _line_size_list(para) -> list:
    """_line_sizes, line by line in order."""
    sizes, line = [], []
    for run in para.runs:
        for k, piece in enumerate(run.text.split("\v")):
            if k and line:
                sizes.append(max(line, key=lambda s: s or 0.0) if None not in line else None)
                line = []
            if piece:
                line.append(_ppt_size(run, para))
    if line:
        sizes.append(max(line, key=lambda s: s or 0.0) if None not in line else None)
    return sizes


def _line_tightening(para, width_pt: float) -> dict:
    """Letter spacing (pt) to take off each line of a paragraph kept on
    Publisher's lines, keyed by line number, for lines that only just fit."""
    lines = [[]]
    for run in para.runs:
        for k, piece in enumerate(run.text.split("\v")):
            if k:
                lines.append([])
            lines[-1].append((run, piece))
    stops = sorted(pos for pos, _ in para.tab_stops)
    # Publisher's own widths, where its lines are these lines, for text it
    # fits to its box: PowerPoint sets some wider than the estimate, as
    # Publisher measures it
    rec = (para.pub or {}) if para.best_fit else {}
    pub_widths = rec.get("lineWidths") or []
    pub_text = str(rec.get("text", ""))
    cuts = [_utf16_index(pub_text, c) for c in rec.get("lineStarts") or []] + [len(pub_text)]
    if len(pub_widths) != len(lines) or len(cuts) != len(lines) + 1:
        pub_widths = []
    out = {}
    for i, pieces in enumerate(lines):
        # where the line ends, with PowerPoint's extra width, following tabs
        # to their stops; only the text after the last tab can be squeezed
        x = para.left_indent_pt + (para.indent_pt if i == 0 else 0.0)
        chars = 0
        # a hanging first line's tab stops at the indent, before the defaults
        line_stops = (sorted(stops + [para.left_indent_pt])
                      if i == 0 and para.indent_pt < 0 else stops)
        for k, (run, piece) in enumerate(pieces):
            if k == len(pieces) - 1:
                piece = piece.rstrip()
            for j, seg in enumerate(piece.split("\t")):
                if j:
                    x = next_tab_px(x, line_stops, TAB_STOP_PT)
                    chars = 0
                seg_w = _run_text_width(run, seg)
                x += (seg_w + _keep_lines_extra(seg_w, len(seg), run.size_pt or 12.0, run.font)
                      + (run.spacing_pt or 0.0) * len(seg))
                chars += len(seg)
        pub_line = pub_text[cuts[i]:cuts[i + 1]].rstrip("\r\v\n") if pub_widths else ""
        trail = pub_line[len(pub_line.rstrip(" \xa0")):]
        # a line's bounds take in the space it wraps at (and run over a run
        # of spaces out to the frame's edge)
        if pub_line.strip() and "\t" not in pub_line and len(trail) < 2 and pieces:
            run = pieces[-1][0]
            pub_w = float(pub_widths[i]) - text_width_px(trail, run.font, run.size_pt or 12.0,
                                                         bool(run.bold), bool(run.italic))
            lead = para.left_indent_pt + (para.indent_pt if i == 0 else 0.0)
            size = max((run.size_pt or 12.0) for run, _ in pieces)
            x = max(x, lead + pub_w + _keep_lines_extra(pub_w, len(pub_line.strip()), size, run.font))
        need = x - width_pt
        if need > 0 and chars:
            size = max((run.size_pt or 12.0) for run, _ in pieces)
            # PowerPoint's letter spacing is in hundredths of a point
            out[i] = min(math.ceil(need / chars * 100) / 100, KEEP_LINES_MAX_SPC * size)
    return out


def scale_slide(slide, s: float):
    """Shrink everything on a slide by s: positions and sizes, and the
    point sizes of text, spacing, indents, tabs and outlines."""
    emu_attrs = {qn("a:off"): ("x", "y"), qn("a:ext"): ("cx", "cy"),
                 qn("a:pPr"): ("marL", "marR", "indent", "defTabSz"),
                 qn("a:tab"): ("pos",), qn("a:ln"): ("w",), qn("a:bodyPr"): ("spcCol",)}
    for el in slide.shapes._spTree.iter():
        names = emu_attrs.get(el.tag, ())
        if el.tag in (qn("a:rPr"), qn("a:endParaRPr")):
            names = ("sz", "spc")
        elif el.tag == qn("a:spcPts"):
            names = ("val",)
        for name in names:
            v = el.get(name)
            if v is not None:
                v = int(round(int(v) * s))
                # PowerPoint's smallest font size is 1pt
                el.set(name, str(max(v, 100) if name == "sz" else v))


def build_slide(prs, boxes, resolver, dpi, warn, verbose, scale=1.0):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # the page before scale_slide shrinks it to fit
    sw_px = prs.slide_width / EMU_PER_INCH * dpi / scale
    sh_px = prs.slide_height / EMU_PER_INCH * dpi / scale

    for b in sorted(boxes, key=lambda z: (z.z or 0, z.order)):
        w = max(b.w or 1.0, 1.0)
        h = max(b.h or 1.0, 1.0)
        # keep shapes from running miles off the canvas
        if b.x > sw_px * 1.5 or b.y > sh_px * 1.5:
            warn(f"shape far outside the page, kept anyway ({int(b.x)},{int(b.y)}px)")
        left, top = px_to_emu(b.x, dpi), px_to_emu(b.y, dpi)
        width, height = px_to_emu(w, dpi), px_to_emu(h, dpi)

        if b.kind == "image":
            path = resolver.resolve(b.src)
            if not path:
                continue
            try:
                pic = slide.shapes.add_picture(path, left, top, width, height)
            except Exception as exc:                   # noqa: BLE001
                warn(f"could not place image {os.path.basename(path)}: {exc}")
                continue
            if b.geom in SHAPE_GEOMS and b.geom != "rect":
                # cropped to a shape, as in Publisher
                pic.auto_shape_type = SHAPE_GEOMS[b.geom]
                av = pic._element.spPr.find(qn("a:prstGeom")).find(qn("a:avLst"))
                for i, v in enumerate(b.adj):
                    av.append(av.makeelement(qn("a:gd"), {
                        "name": "adj" if i == 0 else f"adj{i + 1}",
                        "fmla": f"val {int(round(v * 100000))}"}))
            continue

        if b.kind == "rect":
            shp = slide.shapes.add_shape(SHAPE_GEOMS.get(b.geom, MSO_SHAPE.RECTANGLE),
                                         left, top, width, height)
            for i, v in enumerate(b.adj[:len(shp.adjustments)]):
                shp.adjustments[i] = v
            if b.fill is not None:
                shp.fill.solid()
                shp.fill.fore_color.rgb = b.fill
            else:
                shp.fill.background()
            if b.line is not None:
                shp.line.color.rgb = b.line
                shp.line.width = Pt(max(b.line_w_px * 72.0 / dpi, 0.5))
                if b.dash in LINE_DASHES:
                    shp.line.dash_style = LINE_DASHES[b.dash]
                if b.round_cap:
                    shp.line._get_or_add_ln().set("cap", "rnd")
            else:
                shp.line.fill.background()
            shp.shadow.inherit = False
            continue

        # text
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"middle": MSO_ANCHOR.MIDDLE,
                              "bottom": MSO_ANCHOR.BOTTOM}.get(b.anchor, MSO_ANCHOR.TOP)
        # the width a line has: a column's
        line_w = w
        if b.columns > 1:
            body = tf._txBody.bodyPr
            body.set("numCol", str(b.columns))
            body.set("spcCol", str(px_to_emu(b.column_gap, dpi)))
            line_w = (w - (b.columns - 1) * b.column_gap) / b.columns
        # PowerPoint rounds spacing in points to whole points, so each
        # paragraph's rounding is carried into the next one's space after
        carry = 0.0
        for pi, para in enumerate(b.paras):
            p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
            if para.align is not None:
                p.alignment = para.align
            sizes = _line_sizes(para)
            if para.line_multiple:
                p.line_spacing = para.line_multiple
            elif para.exact_line_pt and len(sizes) == 1 and None not in sizes:
                # a multiple isn't rounded: PowerPoint's single spacing is
                # 1.2 times the largest font size on the line, whatever the font
                p.line_spacing = para.exact_line_pt / (1.2 * sizes.pop())
            elif para.exact_line_pt:
                p.line_spacing = Pt(para.exact_line_pt)
            elif para.line_spacing:
                p.line_spacing = para.line_spacing
            if "\t" in para.text():
                ppr = p._p.get_or_add_pPr()
                ppr.set("defTabSz", str(int(Pt(TAB_STOP_PT))))
                if para.tab_stops:
                    tab_lst = ppr.makeelement(qn("a:tabLst"), {})
                    for pos_pt, algn in sorted(para.tab_stops):
                        tab_lst.append(ppr.makeelement(qn("a:tab"), {
                            "pos": str(int(Pt(pos_pt))), "algn": TAB_ALIGN.get(algn, "l")}))
                    # schema order: after spacing, which is already set
                    ppr.append(tab_lst)
            if para.left_indent_pt:
                p._p.get_or_add_pPr().set("marL", str(int(Pt(min(para.left_indent_pt, 288)))))
            if para.right_indent_pt:
                p._p.get_or_add_pPr().set("marR", str(int(Pt(min(para.right_indent_pt, 288)))))
            if para.indent_pt:
                # first-line indent, negative for a hanging one: python-pptx
                # has no API for either
                p._p.get_or_add_pPr().set("indent", str(int(Pt(
                    max(min(para.indent_pt, 144), -para.left_indent_pt)))))
            if para.space_after_pt or carry:
                # the HTML's margins can be wild, but a gap Publisher placed
                # is real, up to PowerPoint's limit of 1584pt
                want = min(para.space_after_pt, 1584 if para.pub_placed else 48) + carry
                after = max(round(want), 0)
                carry = want - after
                if after:
                    p.space_after = Pt(after)
            tighten = (_line_tightening(para, line_w * 72.0 / dpi - para.right_indent_pt)
                       if para.keep_lines else {})
            line = 0
            for run in para.runs:
                # "\v" is a line break inside the paragraph
                for k, piece in enumerate(run.text.split("\v")):
                    if k:
                        line += 1
                        p.add_line_break()
                        if run.size_pt:
                            # sized like its line: a bare break is 18pt
                            p._p[-1].get_or_add_rPr().set("sz", str(int(run.size_pt * 100)))
                    if not piece:
                        continue
                    r = p.add_run()
                    r.text = piece
                    f = r.font
                    if run.font:
                        f.name = run.font
                    # (the export gives a superscript's shrunk size, which
                    # PowerPoint would shrink again)
                    f.size = Pt(max(min(_ppt_size(run, para) or 12.0, 400.0), 1.0))
                    f.bold = run.bold
                    f.italic = run.italic
                    f.underline = run.underline
                    if run.color is not None:
                        f.color.rgb = run.color
                    rpr = r._r.get_or_add_rPr()
                    if run.caps:
                        rpr.set("cap", run.caps)
                    if run.script:
                        rpr.set("baseline", SUPER_BASELINE if run.script == "super"
                                else SUB_BASELINE)
                    spacing = (run.spacing_pt or 0.0) - tighten.get(line, 0.0)
                    if round(spacing * 100):
                        rpr.set("spc", str(int(round(max(min(spacing, 100.0), -100.0) * 100))))
        if verbose:
            print(f"    text  {int(b.x):>5},{int(b.y):>5}  {int(w):>4}x{int(h):<4}  "
                  f"{b.paras[0].text()[:40]!r}" if b.paras else "")
    if scale != 1.0:
        scale_slide(slide, scale)
    return slide


def add_image_slide(prs, path, dpi, caption=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sw, sh = prs.slide_width, prs.slide_height
    ar = aspect_of(path)
    margin = int(0.4 * EMU_PER_INCH)
    avail_w, avail_h = sw - 2 * margin, sh - 2 * margin
    if ar:
        if avail_w / avail_h > ar:
            h = avail_h
            w = int(h * ar)
        else:
            w = avail_w
            h = int(w / ar)
    else:
        w, h = avail_w, avail_h
    slide.shapes.add_picture(path, int((sw - w) / 2), int((sh - h) / 2), w, h)
    if caption:
        tb = slide.shapes.add_textbox(margin, int(sh - 0.38 * EMU_PER_INCH),
                                      sw - 2 * margin, int(0.28 * EMU_PER_INCH))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = caption
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string("808080")
    return slide


def add_text_slide(prs, title, body):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    margin = int(0.5 * EMU_PER_INCH)
    tb = slide.shapes.add_textbox(margin, margin,
                                  prs.slide_width - 2 * margin,
                                  prs.slide_height - 2 * margin)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(16)
    r.font.bold = True
    for line in body.splitlines():
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(11)
    return slide


# ----------------------------------------------------------------------------
# Driver
# ----------------------------------------------------------------------------

def natural_key(path: str):
    name = os.path.basename(path).lower()
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", name)]


def html_files_for(export_dir: str, base: str) -> list:
    """Publisher writes <base>.htm for page 1 plus page*.htm for the rest."""
    found = []
    for pattern in ("*.htm", "*.html"):
        found.extend(globmod.glob(os.path.join(export_dir, pattern)))
        found.extend(globmod.glob(os.path.join(export_dir, f"{base}_files", pattern)))
    seen, out = set(), []
    primary = os.path.join(export_dir, base + ".htm")
    if os.path.isfile(primary):
        out.append(primary)
        seen.add(os.path.normcase(primary))
    for f in sorted(found, key=natural_key):
        nc = os.path.normcase(f)
        if nc in seen:
            continue
        if os.path.basename(f).lower() in ("index.htm", "index.html") and out:
            continue
        seen.add(nc)
        out.append(f)
    return out


def convert_export(export_dir: str, out_path: str, args) -> dict:
    base = os.path.basename(os.path.normpath(export_dir))
    warnings: list = []

    def warn(msg):
        warnings.append(msg)
        if args.verbose:
            print(f"    ! {msg}")

    htmls = html_files_for(export_dir, base)
    if not htmls:
        raise FileNotFoundError(f"no .htm found in {export_dir}")

    fallback = tuple(float(v) for v in args.page_size.lower().split("x"))
    # Publisher sizes its page wrapper to the content, not the page, and
    # sometimes writes a malformed height, so prefer stage 1's real size.
    real_page = read_page_size(export_dir, base)
    dpi = float(args.dpi)
    resolver = ImageResolver(export_dir, export_dir, warn)
    conv = Converter(args, warn)
    conv.pub_text = load_pub_text(export_dir, base)
    conv.pub_blank_frames = load_pub_frames(export_dir, base)
    conv.pub_tables = load_pub_tables(export_dir, base)
    conv.pub_inline = load_pub_inline(export_dir, base)
    check_pub_inline(export_dir, base, real_page, warn)
    pub_pictures = load_pub_pictures(export_dir, base)

    prs = Presentation()
    slide_pages = []   # (page_w_px, page_h_px, boxes)

    for html_path in htmls:
        with open(html_path, "rb") as fh:
            raw = fh.read()
        soup = BeautifulSoup(drop_downlevel_fallbacks(raw), "lxml")
        mark_inset_pens(soup)
        restore_vml_text_boxes(soup)
        restore_vml_groups(soup, pub_pictures, len(slide_pages) + 1)
        conv.inline_boxes = restore_vml_inline_text_boxes(soup)
        restore_vml_pictures(soup)
        conv.vml_order = vml_shape_order(soup)
        sheet = StyleSheet()
        for st in soup.find_all("style"):
            sheet.add_css(st.get_text())
        for link in soup.find_all("link", rel=lambda v: v and "stylesheet" in " ".join(v).lower()):
            href = unquote(link.get("href") or "")
            cand = os.path.normpath(os.path.join(os.path.dirname(html_path), href))
            if os.path.isfile(cand):
                try:
                    with open(cand, "rb") as fh:
                        sheet.add_css(fh.read().decode("utf-8", "replace"))
                except OSError:
                    warn(f"could not read stylesheet {href}")

        body = soup.body or soup
        body_style = computed_style(body, sheet, {}) if isinstance(body, Tag) else {}
        pages = find_page_containers(soup, sheet, dpi,
                                     (real_page or fallback)[1] * dpi,
                                     (real_page or fallback)[0] * dpi)

        resolver.html_dir = os.path.dirname(html_path)

        if pages:
            for el, w, h, st in pages:
                conv.order = 0
                conv.page = len(slide_pages) + 1
                boxes = []
                conv.walk(el, sheet, {**body_style, **st}, (0.0, 0.0), w, h, boxes)
                if real_page:
                    w, h = real_page[0] * dpi, real_page[1] * dpi
                else:
                    # the wrapper can be smaller than its content
                    w, h = bounding_page(boxes, dpi, (w / dpi, h / dpi))
                slide_pages.append((w, h, boxes))
        else:
            conv.order = 0
            conv.page = len(slide_pages) + 1
            boxes = []
            page_in = real_page or fallback
            pw = page_in[0] * dpi
            ph = page_in[1] * dpi
            conv.walk(body, sheet, body_style, (0.0, 0.0), pw, ph, boxes)
            if not real_page:
                pw, ph = bounding_page(boxes, dpi, fallback)
            slide_pages.append((pw, ph, boxes))

    if not slide_pages:
        raise ValueError("no page content found in the HTML")

    # one deck-wide slide size: PowerPoint/Canva allow only one
    page_w = max(p[0] for p in slide_pages)
    page_h = max(p[1] for p in slide_pages)
    w_in = max(page_w / dpi, 1.0)
    h_in = max(page_h / dpi, 1.0)
    # a page larger than PowerPoint's 56in limit (a banner) is shrunk to
    # fit, keeping its shape, rather than cut off
    scale = min(1.0, PPTX_MAX_IN / w_in, PPTX_MAX_IN / h_in)
    if scale < 1.0:
        warn(f"page is {w_in:.2f} x {h_in:.2f} in, larger than PowerPoint's "
             f"{PPTX_MAX_IN:.0f} in limit; the slide is scaled to {scale:.0%}")
        w_in, h_in = min(w_in * scale, PPTX_MAX_IN), min(h_in * scale, PPTX_MAX_IN)
    prs.slide_width = Emu(int(round(w_in * EMU_PER_INCH)))
    prs.slide_height = Emu(int(round(h_in * EMU_PER_INCH)))
    sizes = {(round(p[0]), round(p[1])) for p in slide_pages}
    if len(sizes) > 1:
        warn(f"pages have {len(sizes)} different sizes; the deck uses the largest "
             f"({w_in:.2f} x {h_in:.2f} in) and smaller pages sit top-left")

    if args.hires:
        for idx, (pw, ph, boxes) in enumerate(slide_pages, start=1):
            apply_hires(boxes, export_dir, base, idx, resolver, warn, dpi)
    # after apply_hires, which matches pictures to the page they belong to
    spreads = read_page_spreads(export_dir, base)
    if spreads and len(spreads) != len(slide_pages):
        warn(f"Publisher has {len(spreads)} pages but the HTML has {len(slide_pages)}; "
             f"shapes crossing a spread's fold are left on their own page")
    elif spreads:
        shared = share_spread_shapes(slide_pages, spreads, dpi)
        if shared and args.verbose:
            print(f"  {shared} shape(s) crossing a fold copied to the facing page")

    for idx, (pw, ph, boxes) in enumerate(slide_pages, start=1):
        if args.verbose:
            print(f"  page {idx}: {len(boxes)} shapes  ({pw/dpi:.2f} x {ph/dpi:.2f} in)")
        build_slide(prs, boxes, resolver, dpi, warn, args.verbose, scale)

    extra = 0
    if not args.no_extras:
        extra = append_extras(prs, export_dir, base, resolver, dpi, warn)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    prs.save(out_path)

    stats = {
        "pptx": out_path,
        "slides": len(slide_pages) + extra,
        "pages": len(slide_pages),
        "extras": extra,
        "shapes": sum(len(p[2]) for p in slide_pages),
        "warnings": warnings,
        "size_in": (w_in, h_in),
    }

    if args.report:
        report = os.path.splitext(out_path)[0] + "_conversion_report.txt"
        with open(report, "w", encoding="utf-8") as fh:
            fh.write(f"Source export folder : {export_dir}\n")
            fh.write(f"PowerPoint written   : {out_path}\n")
            fh.write(f"Slide size           : {w_in:.2f} x {h_in:.2f} in\n")
            fh.write(f"Layout slides        : {len(slide_pages)}\n")
            fh.write(f"Extra slides         : {extra}\n")
            fh.write(f"Shapes placed        : {stats['shapes']}\n")
            fh.write(f"HTML files read      : {len(htmls)}\n\n")
            if warnings:
                fh.write("Warnings:\n")
                for w in warnings:
                    fh.write(f"  - {w}\n")
            else:
                fh.write("No warnings.\n")
        stats["report"] = report
    return stats


def apply_hires(boxes, export_dir, base, page_no, resolver, warn, dpi):
    """Swap Publisher's web-quality images for the 300dpi COM PNGs: each for
    the image in the same place, where stage 1 recorded the places, else
    in order."""
    img_boxes = [b for b in sorted(boxes, key=lambda z: z.order) if b.kind == "image"]
    if not img_boxes:
        return
    pics = load_pub_pictures(export_dir, base)
    if pics is not None:
        _hires_by_place(img_boxes, [p for p in pics if p["page"] in (0, page_no)],
                        resolver, dpi)
        return
    cands = hires_candidates(export_dir, base, page_no)
    if not cands:
        return
    if len(cands) != len(img_boxes):
        warn(f"page {page_no}: {len(img_boxes)} images in HTML but {len(cands)} "
             f"exported PNGs; high-res swap skipped for this page")
        return
    for box, png in zip(img_boxes, cands):
        box_ar = (box.w / box.h) if box.h else None
        png_ar = aspect_of(png)
        if box_ar and png_ar and abs(box_ar - png_ar) / max(box_ar, png_ar) > 0.12:
            warn(f"page {page_no}: aspect mismatch for {os.path.basename(png)}; kept web image")
            continue
        box.src = os.path.basename(png)
        resolver.used.add(os.path.normcase(os.path.abspath(png)))


def _place_misfit(x, y, w, h, pic, k=1.0) -> Optional[float]:
    """How far a frame at x, y, w, h (pt times k) is from stage 1's picture,
    by centre and by size either way round, since the HTML gives a rotated
    picture's frame unrotated or its bounds; None if too far to be it."""
    pw, ph = pic["width"] * k, pic["height"] * k
    if not w or not h or pw <= 0 or ph <= 0:
        return None
    off = math.hypot(x + w / 2 - (pic["left"] * k + pw / 2),
                     y + h / 2 - (pic["top"] * k + ph / 2))
    if off > max(3.0 * k, 0.08 * min(pw, ph)):
        return None
    size = min(max(abs(w - pw) / pw, abs(h - ph) / ph),
               max(abs(w - ph) / ph, abs(h - pw) / pw))
    if size > 0.12:
        return None
    return off / max(pw, ph) + size


def _hires_by_place(img_boxes, pics, resolver, dpi):
    """Gives each image the PNG of the picture stage 1 saw in its place. An
    image with none, a picture of WordArt or a shape, keeps the web copy."""
    k = dpi / 72.0
    pairs = []
    for bi, box in enumerate(img_boxes):
        for pi, pic in enumerate(pics):
            misfit = _place_misfit(box.x, box.y, box.w, box.h, pic, k)
            if misfit is not None:
                pairs.append((misfit, bi, pi))
    used_b, used_p = set(), set()
    for _, bi, pi in sorted(pairs):
        if bi in used_b or pi in used_p:
            continue
        png = pics[pi]["png"]
        # a PNG of another shape than the image's, a cropped picture's
        # whole one for instance, would stretch
        box_ar, png_ar = img_boxes[bi].w / img_boxes[bi].h, aspect_of(png)
        if png_ar and min(abs(box_ar - png_ar) / max(box_ar, png_ar),
                          abs(1 / box_ar - png_ar) / max(1 / box_ar, png_ar)) > 0.12:
            continue
        used_b.add(bi)
        used_p.add(pi)
        img_boxes[bi].src = os.path.basename(png)
        resolver.used.add(os.path.normcase(os.path.abspath(png)))


def append_extras(prs, export_dir, base, resolver, dpi, warn) -> int:
    """Scratch-area art and text never reach the HTML; carry them in on their own slides."""
    added = 0
    patterns = [f"{base}_scratch_*_image_*.png", f"{base}_scratch_*_text_*.png"]
    files = []
    for pat in patterns:
        files.extend(sorted(globmod.glob(os.path.join(export_dir, pat)), key=natural_key))
    for f in files:
        if os.path.normcase(os.path.abspath(f)) in resolver.used:
            continue
        try:
            add_image_slide(prs, f, dpi, caption=f"Scratch area (off-canvas): {os.path.basename(f)}")
            added += 1
        except Exception as exc:                       # noqa: BLE001
            warn(f"could not add scratch image {os.path.basename(f)}: {exc}")

    txt = os.path.join(export_dir, base + "_scratch_text.txt")
    if os.path.isfile(txt):
        try:
            with open(txt, "r", encoding="utf-8", errors="replace") as fh:
                content = fh.read().strip()
            if content:
                lines = content.splitlines()
                per = 34
                for i in range(0, len(lines), per):
                    add_text_slide(prs, "Scratch-area text (not in the HTML export)",
                                   "\n".join(lines[i:i + per]))
                    added += 1
        except OSError as exc:
            warn(f"could not read {os.path.basename(txt)}: {exc}")
    return added


def scan_tree(root: str) -> dict:
    """Walk the tree once and report what is actually there.

    An export folder is any folder holding at least one .htm/.html file. The
    earlier rule demanded <FolderName>.htm, which missed exports where Publisher
    named the page differently (index.htm and friends) or where the folder had
    been renamed after the export.
    """
    export_dirs, pub_files, companion_only = [], [], []
    for dirpath, dirnames, filenames in os.walk(root):
        base = os.path.basename(os.path.normpath(dirpath))

        # Publisher's own asset folder: its pages are read via the parent export
        # folder, so it is never an export root in its own right.
        dirnames[:] = [d for d in dirnames if not d.lower().endswith("_files")]
        if base.lower().endswith("_files"):
            continue

        pub_files.extend(os.path.join(dirpath, f) for f in filenames
                         if f.lower().endswith(".pub"))

        htms = [f for f in filenames if f.lower().endswith((".htm", ".html"))]
        if htms:
            export_dirs.append(dirpath)
        elif any(os.path.isdir(os.path.join(dirpath, d)) and d.lower().endswith("_files")
                 for d in os.listdir(dirpath) if os.path.isdir(os.path.join(dirpath, d))):
            companion_only.append(dirpath)

    return {
        "exports": sorted(set(export_dirs)),
        "pubs": sorted(pub_files),
        "companion_only": sorted(companion_only),
    }


def find_export_dirs(root: str) -> list:
    return scan_tree(root)["exports"]


def explain_no_exports(root: str, scan: dict) -> str:
    """Say what was found instead, so the failure diagnoses itself."""
    lines = [f"No Publisher export folders found beneath {root}"]
    lines.append("  (an export folder is one containing a .htm or .html file)")

    pubs = scan["pubs"]
    if pubs:
        lines.append("")
        lines.append(f"  Found {len(pubs)} Publisher file(s) here but no HTML export beside them:")
        for p in pubs[:5]:
            lines.append(f"    {p}")
        if len(pubs) > 5:
            lines.append(f"    ... and {len(pubs) - 5} more")
        lines.append("")
        lines.append("  Run the export script first, in the folder holding the .pub files:")
        lines.append('    .\\PublisherPubToHtmlForPptx.ps1 -Filter "*.pub" -Recurse')
        lines.append("")
        lines.append("  That leaves one subfolder per publication containing <BaseName>.htm,")
        lines.append("  which is what this converter reads.")
    elif scan["companion_only"]:
        lines.append("")
        lines.append("  Found Publisher _files asset folder(s) with no .htm beside them:")
        for d in scan["companion_only"][:5]:
            lines.append(f"    {d}")
        lines.append("")
        lines.append("  The .htm was probably moved or deleted. Re-run the export script,")
        lines.append("  or delete the export subfolder so it will be rebuilt (the export")
        lines.append("  script skips a publication whose .htm already exists).")
    else:
        lines.append("")
        lines.append("  Nothing that looks like a Publisher export is in that tree.")
        lines.append("  Check the path, and that you pointed at the folder holding the")
        lines.append("  .pub files or their export subfolders.")
    return "\n".join(lines)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="Convert Publisher filtered-HTML exports into layout-preserving .pptx for Canva.")
    ap.add_argument("input", help="export folder, a .htm file, or a tree with --recurse")
    ap.add_argument("-o", "--output", help="output .pptx, or output folder with --recurse")
    ap.add_argument("--recurse", action="store_true", help="process every export folder below input")
    ap.add_argument("--dpi", type=float, default=96.0, help="CSS pixels per inch (default 96)")
    ap.add_argument("--hires", action="store_true", help="use the 300dpi COM-exported PNGs")
    ap.add_argument("--no-extras", action="store_true", help="skip scratch-area slides")
    ap.add_argument("--reflow", action="store_true",
                    help="let text wrap freely instead of breaking lines where Publisher does")
    ap.add_argument("--page-size", default="8.5x11", help="fallback page size in inches")
    ap.add_argument("--report", action="store_true", help="write a per-file conversion report")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args(argv)

    inp = os.path.abspath(args.input)
    if not os.path.exists(inp):
        print(f"Input not found: {inp}", file=sys.stderr)
        return 1

    if os.path.isfile(inp):
        targets = [os.path.dirname(inp)]
    elif args.recurse:
        scan = scan_tree(inp)
        targets = scan["exports"]
        if not targets:
            print(explain_no_exports(inp, scan), file=sys.stderr)
            return 1
    else:
        targets = [inp]

    print("")
    print("Publisher HTML export -> PowerPoint (Canva-ready)")
    print("=" * 52)

    ok = failed = 0
    total_warn = 0
    for d in targets:
        base = os.path.basename(os.path.normpath(d))
        if args.output and not args.recurse and os.path.splitext(args.output)[1].lower() == ".pptx":
            out = os.path.abspath(args.output)
        else:
            out_dir = os.path.abspath(args.output) if args.output else d
            out = os.path.join(out_dir, base + ".pptx")
        print(f"\n{base}")
        try:
            stats = convert_export(d, out, args)
        except Exception as exc:                        # noqa: BLE001
            print(f"  FAILED: {exc}")
            if args.verbose:
                import traceback
                traceback.print_exc()
            failed += 1
            continue
        ok += 1
        total_warn += len(stats["warnings"])
        print(f"  {stats['pages']} page slide(s) + {stats['extras']} extra, "
              f"{stats['shapes']} shapes, {stats['size_in'][0]:.2f} x {stats['size_in'][1]:.2f} in")
        print(f"  -> {out}")
        if stats["warnings"] and not args.verbose:
            print(f"  {len(stats['warnings'])} warning(s)"
                  + (f"; see {os.path.basename(stats['report'])}" if args.report else
                     " (re-run with --report or -v for detail)"))

    print("")
    print("=" * 52)
    print(f"Converted: {ok}    Failed: {failed}    Warnings: {total_warn}")
    print("")
    print("In Canva: Create a design > Import file, or drag the .pptx onto the")
    print("Projects page. Fonts Canva does not have are substituted, so check")
    print("line breaks on each page before sending anything to print.")
    print("")
    return 0 if failed == 0 else 2


if __name__ == "__main__":
    sys.exit(main())